"""Text extraction for uploaded documents: PDF, DOCX, TXT/MD/others treated as plain text.

Returns pages: [{page: int, text: str}] so citations can point at a page number.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any

DOC_EXTS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".rtf", ".csv", ".json", ".html", ".htm", ".srt", ".vtt"}
MEDIA_EXTS = {".mp3", ".m4a", ".wav", ".aac", ".flac", ".ogg", ".opus", ".mp4", ".mov", ".m4v", ".mkv", ".webm", ".avi", ".aiff"}


def is_document(path: Path) -> bool:
    return path.suffix.lower() in DOC_EXTS


def is_media(path: Path) -> bool:
    return path.suffix.lower() in MEDIA_EXTS


def _reflow(text: str) -> str:
    """Some PDF exporters (Google Docs among them) emit one word per line, which turns a page into a column of
    single words: retrieval and findings then see no phrases at all. When a page's lines average fewer than
    REFLOW_MAX_WORDS_PER_LINE words, join them back into running text (blank lines still separate paragraphs).
    Ordinary pages — sentences per line — are left exactly as extracted."""
    lines = [ln.rstrip() for ln in text.splitlines()]
    words = [len(ln.split()) for ln in lines if ln.strip()]
    if len(words) < 8 or sum(words) / len(words) >= REFLOW_MAX_WORDS_PER_LINE:
        return text
    out: list[str] = []
    para: list[str] = []
    for ln in lines:
        if ln.strip():
            para.append(ln.strip())
        elif para:
            out.append(" ".join(para))
            para = []
    if para:
        out.append(" ".join(para))
    return "\n\n".join(out)


REFLOW_MAX_WORDS_PER_LINE = 2.5


def extract_pages(path: Path) -> list[dict[str, Any]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text() or ""
            except Exception:  # noqa: BLE001
                text = ""
            if text.strip():
                pages.append({"page": i, "text": _reflow(text)})
        if not pages:
            raise RuntimeError("no extractable text in PDF (scanned? run OCR first)")
        return pages
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        blocks = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                blocks.append(" | ".join(c.text for c in row.cells))
        text = "\n\n".join(b for b in blocks if b.strip())
        return _paginate(text)
    if ext in (".html", ".htm"):
        import re

        raw = path.read_text(errors="replace")
        raw = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", raw, flags=re.S | re.I)
        raw = re.sub(r"<br\s*/?>|</p>|</div>|</h\d>", "\n\n", raw, flags=re.I)
        text = re.sub(r"<[^>]+>", " ", raw)
        return _paginate(text)
    if ext in (".srt", ".vtt"):
        # subtitle files are really media transcripts; handled by ingest via media.parse_vtt
        raise RuntimeError("subtitle files are ingested as transcripts")
    return _paginate(path.read_text(errors="replace"))


def _paginate(text: str, chars_per_page: int = 3000) -> list[dict[str, Any]]:
    """Fake page numbers for formats without pages, so citations still have a locator."""
    import re

    paras = [" ".join(p.split()) for p in re.split(r"\n\s*\n", text) if p.strip()]
    pages: list[dict[str, Any]] = []
    cur: list[str] = []
    size = 0
    for p in paras:
        if cur and size + len(p) > chars_per_page:
            pages.append({"page": len(pages) + 1, "text": "\n\n".join(cur)})
            cur, size = [], 0
        cur.append(p)
        size += len(p)
    if cur:
        pages.append({"page": len(pages) + 1, "text": "\n\n".join(cur)})
    if not pages:
        raise RuntimeError("document is empty")
    return pages
