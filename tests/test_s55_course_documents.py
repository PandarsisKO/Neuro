"""S55 — a course's DOCUMENTS come in with its videos (CS7, extension 1.9.6, 2026-09-18). (Sorts after test_core.)

Kyle, on Acquisition Ace: "there are bonus lessons without video content, I still want the information,
spreadsheets, pdfs etc that are in them." Measured live: nine bonus lessons, each a Google Drive share link to a
PDF (a checklist, worksheets, a deal calculator, a playbook), no player, no text of note on the page. The scanner
called them `no_video` and the importer dropped them.

Now: `scan-lib.js` records each lesson's `attachments` (stable document identities: Drive/Docs/Sheets/Slides ids,
Dropbox paths, direct document addresses) and a lesson that is only a link is `document_found`, a ready outcome;
`courses.import_course` queues each distinct document once through `ingest_url`; `ingest.ingest_document_url`
fetches the file through the network boundary (no cookies — a document host is not a video host) and reads it by
exactly the path an uploaded file takes (PDF/DOCX → document, XLSX/CSV → spreadsheet + calculator), remembering
the link on the source so a re-import is answered from the library. A share link that answers with a sign-in
page is reported as blocked, never read as the document.
"""
from __future__ import annotations

import io
import json
import shutil
import subprocess
from pathlib import Path

import pytest

from neurosearch import db

H = {"Authorization": "Bearer t0k"}
ROOT = Path(__file__).resolve().parent.parent
EXT = ROOT / "extension"


def _pdf_bytes(text: str) -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter()
    page = w.add_blank_page(width=300, height=300)
    # a real text object so extract_text() finds something: pypdf writes content streams verbatim
    from pypdf.generic import DecodedStreamObject, NameObject, DictionaryObject
    stream = DecodedStreamObject()
    stream.set_data(f"BT /F1 12 Tf 20 150 Td ({text}) Tj ET".encode("latin-1"))
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): w._add_object(font)})})
    page[NameObject("/Contents")] = w._add_object(stream)
    buf = io.BytesIO(); w.write(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Deal"
    ws["A1"], ws["B1"] = "Purchase price", 1000000
    ws["A2"], ws["B2"] = "SDE", 250000
    ws["A3"], ws["B3"] = "Multiple", "=B1/B2"
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ the link → the download address

def test_share_links_resolve_to_their_download_address_and_pages_do_not():
    from neurosearch.webpage import document_download_url, looks_like_document
    assert document_download_url("https://drive.google.com/file/d/1zgpfCLYr-EOKAgGA3M80wugYEMEbfS1z/view?usp=sharing") == \
        "https://drive.google.com/uc?export=download&id=1zgpfCLYr-EOKAgGA3M80wugYEMEbfS1z"
    assert document_download_url("https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOp/edit#gid=0").endswith("/export?format=xlsx")
    assert document_download_url("https://docs.google.com/document/d/1AbCdEfGhIjKlMnOp/edit").endswith("/export?format=pdf")
    assert document_download_url("https://www.dropbox.com/scl/fi/abc/Deal.xlsx?rlkey=k&dl=0") == "https://www.dropbox.com/scl/fi/abc/Deal.xlsx?rlkey=k&dl=1"
    assert document_download_url("https://school.example.com/files/worksheet.pdf") == "https://school.example.com/files/worksheet.pdf"
    for page in ("https://school.example.com/lessons/1", "https://drive.google.com/drive/folders/1abc", "https://www.loom.com/share/abc"):
        assert not looks_like_document(page), page


# ------------------------------------------------------------------ the file → a source, by the upload path

def _serve(monkeypatch, ctype: str, body: bytes, headers: dict | None = None, final: str | None = None):
    from neurosearch import webpage
    calls = []
    def fake(url, timeout=60.0, content_class=None):
        calls.append((url, content_class))
        return final or url, ctype, body, dict(headers or {})
    monkeypatch.setattr(webpage, "fetch_with_headers", fake)
    return calls


def test_a_drive_pdf_link_becomes_a_document_source_named_after_its_lesson(client, monkeypatch):
    from neurosearch import ingest
    calls = _serve(monkeypatch, "application/pdf", _pdf_bytes("Fast Track Acquisition Checklist step one call the broker"),
                   headers={"Content-Disposition": 'attachment; filename="Fast Track Acquisition.pdf"'})
    p = client.post("/api/projects", headers=H, json={"name": "Docs course", "brief": "b"}).json()
    link = "https://drive.google.com/file/d/1FASTTRACKFASTTRACKFASTTRACK/view?usp=sharing"
    r = ingest.ingest_url(link, project_id=p["id"], title="Chapter 05 › [Bonus #1] Fast Track Acquisition Checklist")
    assert calls == [("https://drive.google.com/uc?export=download&id=1FASTTRACKFASTTRACKFASTTRACK", "document")], "the DOWNLOAD form is fetched, under document limits"
    assert r["transcript"] == "document" and r["segments"] >= 1
    src = client.get(f"/api/sources/{r['source_id']}", headers=H).json()
    assert src["platform"] == "document" and src["status"] == "ready"
    assert src["title"].startswith("Chapter 05 › [Bonus #1]")
    # the LINK is remembered on the file source, so the library answers a re-import
    assert db.sources_for_urls([link]) == {link: r["source_id"]}
    hits = client.get("/api/search?q=call+the+broker&project_id=" + p["id"], headers=H).json()
    assert hits and hits[0]["source_id"] == r["source_id"]


def test_a_spreadsheet_link_becomes_a_spreadsheet_source_with_its_calculator(client, monkeypatch):
    from neurosearch import ingest
    pytest.importorskip("openpyxl")
    _serve(monkeypatch, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", _xlsx_bytes(),
           headers={"content-disposition": "attachment; filename*=UTF-8''AA%20Deal%20Calculator.xlsx"})
    p = client.post("/api/projects", headers=H, json={"name": "Sheet course", "brief": "b"}).json()
    r = ingest.ingest_url("https://docs.google.com/spreadsheets/d/1DEALCALCDEALCALC/edit", project_id=p["id"], title="Deal Calculator")
    assert r["transcript"] == "spreadsheet" and r["inputs"] >= 1 and r["outputs"] >= 1
    src = client.get(f"/api/sources/{r['source_id']}", headers=H).json()
    assert src["platform"] == "spreadsheet" and src["status"] == "ready"


def test_a_share_link_that_answers_with_a_sign_in_page_is_blocked_not_read(client, monkeypatch):
    from neurosearch import ingest
    from neurosearch.acquire import AcquisitionFailure
    wall = b"<html><head><title>Sign in - Google Accounts</title></head><body>Sign in to continue to Google Drive</body></html>"
    _serve(monkeypatch, "text/html; charset=utf-8", wall, final="https://accounts.google.com/v3/signin/identifier")
    p = client.post("/api/projects", headers=H, json={"name": "Private docs", "brief": "b"}).json()
    with pytest.raises(AcquisitionFailure) as ei:
        ingest.ingest_url("https://drive.google.com/file/d/1PRIVATEPRIVATEPRIVATE/view", project_id=p["id"], title="Private worksheet")
    assert "sign-in" in str(ei.value) and "Upload" in str(ei.value)
    assert db.project_source_ids(p["id"], ready_only=True) == [], "nothing was read as a document"


def test_drives_virus_scan_interstitial_is_followed_once_to_the_file(client, monkeypatch):
    """A larger public Drive file answers `uc?export=download` with an HTML "download anyway" form instead of the
    bytes; the form's own action + hidden inputs (id/export/confirm/uuid) name the real download. Followed once."""
    from neurosearch import ingest, webpage
    interstitial = (b'<html><body><form id="download-form" action="https://drive.usercontent.google.com/download" method="get">'
                    b'<input type="hidden" name="id" value="1BIGFILEBIGFILEBIGFILE"><input type="hidden" name="export" value="download">'
                    b'<input type="hidden" name="confirm" value="t"><input type="hidden" name="uuid" value="abc-123"></form></body></html>')
    pdf = _pdf_bytes("the big playbook")
    calls = []
    def fake(url, timeout=60.0, content_class=None):
        calls.append(url)
        if "drive.usercontent.google.com/download" in url:
            return url, "application/pdf", pdf, {"content-disposition": 'attachment; filename="Playbook.pdf"'}
        return "https://drive.google.com/uc?export=download&id=1BIGFILEBIGFILEBIGFILE", "text/html; charset=utf-8", interstitial, {}
    monkeypatch.setattr(webpage, "fetch_with_headers", fake)
    p = client.post("/api/projects", headers=H, json={"name": "Big doc", "brief": "b"}).json()
    r = ingest.ingest_url("https://drive.google.com/file/d/1BIGFILEBIGFILEBIGFILE/view", project_id=p["id"], title="Playbook")
    assert r["transcript"] == "document"
    assert len(calls) == 2 and calls[1] == "https://drive.usercontent.google.com/download?id=1BIGFILEBIGFILEBIGFILE&export=download&confirm=t&uuid=abc-123"


def test_the_same_bytes_behind_two_links_are_one_source(client, monkeypatch):
    from neurosearch import ingest
    _serve(monkeypatch, "application/pdf", _pdf_bytes("one file two links"))
    p = client.post("/api/projects", headers=H, json={"name": "Dup docs", "brief": "b"}).json()
    a = ingest.ingest_url("https://drive.google.com/file/d/1SAMEBYTESSAMEBYTES01/view", project_id=p["id"], title="A")
    b = ingest.ingest_url("https://drive.google.com/file/d/1SAMEBYTESSAMEBYTES02/view", project_id=p["id"], title="B")
    assert a["source_id"] == b["source_id"], "identity is the file's content fingerprint, as for an upload"


# ------------------------------------------------------------------ the importer

def test_course_import_queues_each_distinct_document_once_and_never_sends_cookies_for_it(client):
    p = client.post("/api/projects", headers=H, json={"name": "Course with bonuses", "brief": "b"}).json()
    drive = "https://drive.google.com/file/d/1SHAREDWORKSHEET01/view"
    r = client.post(f"/api/projects/{p['id']}/course-import", headers=H, json={
        "course": {"title": "Acquisition Profit System", "url": "https://courses.example.com/courses/aps"},
        "lessons": [
            {"title": "Conclusion", "module": "Chapter 05 · Conclusion", "page_url": "https://courses.example.com/courses/aps/conclusion",
             "video_urls": ["https://www.loom.com/embed/conc1"], "outcome": "video_found",
             "attachments": [{"provider": "file", "url": "https://courses.example.com/files/recap.pdf", "title": "recap.pdf", "kind": "pdf"}]},
            {"title": "[Bonus #1] Fast Track Checklist", "module": "Chapter 05 · Conclusion", "page_url": "https://courses.example.com/courses/aps/b1",
             "video_urls": [], "outcome": "document_found", "attachments": [{"provider": "gdrive", "url": drive, "title": "", "kind": "gdrive"}]},
            {"title": "[Bonus #2] Same worksheet again", "module": "Chapter 05 · Conclusion", "page_url": "https://courses.example.com/courses/aps/b2",
             "video_urls": [], "outcome": "document_found", "attachments": [{"provider": "gdrive", "url": drive, "title": "", "kind": "gdrive"}]},
            {"title": "Empty", "module": "Chapter 05 · Conclusion", "page_url": "https://courses.example.com/courses/aps/e", "video_urls": [], "outcome": "no_video", "attachments": []},
        ],
        "cookies": [{"domain": ".courses.example.com", "name": "session", "value": "s3cr3t", "path": "/", "secure": True, "expirationDate": 1900000000}]}).json()
    assert r["queued"] == 1 and r["documents"] == 2, r
    assert r["no_video"] == ["[Bonus #1] Fast Track Checklist", "[Bonus #2] Same worksheet again", "Empty"] and r["nothing"] == ["Empty"]
    jobs_ = [j for j in client.get(f"/api/projects/{p['id']}/jobs", headers=H).json() if j["kind"] == "ingest_url"]
    by_url = {j["payload"]["url"]: j["payload"] for j in jobs_}
    assert set(by_url) == {"https://www.loom.com/share/conc1", "https://courses.example.com/files/recap.pdf", drive}
    assert by_url[drive]["title"] == "Chapter 05 · Conclusion › [Bonus #1] Fast Track Checklist (+1 more lesson)"
    assert by_url[drive]["referer"].endswith("/aps/b1") and "cookies_file" not in by_url[drive], "a document is fetched without the user's session"
    assert by_url["https://courses.example.com/files/recap.pdf"]["title"].endswith("— recap.pdf")
    assert by_url["https://www.loom.com/share/conc1"]["cookies_file"], "the video still gets the session it needs"


# ------------------------------------------------------------------ the extension (contract level; the jsdom traversal is in test_s32)

def _node():
    exe = shutil.which("node")
    if not exe:
        pytest.skip("no JavaScript engine on this machine")
    return exe


def test_document_identity_is_stable_and_a_page_is_not_a_document():
    lib = json.dumps(str(EXT / "scan-lib.js"))
    js = (f"globalThis.document=undefined; eval(require('fs').readFileSync({lib},'utf8'));"
          "const d = NSScan.documentIdentity; console.log(JSON.stringify(["
          "d('https://drive.google.com/file/d/1zgpfCLYr-EOKAgGA3M80wugYEMEbfS1z/view?usp=sharing'),"
          "d('https://drive.google.com/open?id=1zgpfCLYr-EOKAgGA3M80wugYEMEbfS1z'),"
          "d('https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOp/edit#gid=0'),"
          "d('https://www.dropbox.com/scl/fi/abc/Deal.xlsx?dl=0'),"
          "d('https://school.example.com/files/Goal%20Setting.pdf?x=1'),"
          "d('https://school.example.com/lessons/1'), d('https://www.loom.com/share/abc'), d('https://drive.google.com/drive/folders/1abc')]));")
    out = subprocess.run([_node(), "-e", js], capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stderr[:500]
    a, b, c, dbx, f, *pages = json.loads(out.stdout)
    assert a["provider"] == "gdrive" and a["id"] == b["id"] and a["url"] == b["url"] == "https://drive.google.com/file/d/1zgpfCLYr-EOKAgGA3M80wugYEMEbfS1z/view"
    assert c["provider"] == "gsheets" and dbx["provider"] == "dropbox" and f["provider"] == "file" and f["ext"] == "pdf"
    assert pages == [None, None, None]
    assert "document_found" in (EXT / "scan-lib.js").read_text() and "document_found" in (EXT / "background.js").read_text()
    assert "document_found" in (EXT / "popup.js").read_text()
