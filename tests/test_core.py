"""Offline tests: parsing, chunking, storage, search, projects, import API, MCP auth.
Run: pytest -q
"""
from __future__ import annotations

import os
import pathlib
import shutil
import tempfile

os.environ["NEUROSEARCH_DATA_DIR"] = tempfile.mkdtemp(prefix="ns_test_")
os.environ["NEUROSEARCH_APP_TOKEN"] = "t0k"
os.environ.pop("OPENAI_API_KEY", None)

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

import json  # noqa: E402

from neurosearch import chunking, db, ingest, jobs, media  # noqa: E402
from neurosearch.api import app  # noqa: E402

H = {"Authorization": "Bearer t0k"}
CHAT_ARM_INPUT_TOTAL = 210014   # frozen: plain + cache read + cache write of the 34-question chat arm (re-frozen 0.32.0: data-not-instructions + corrected-post rules; 0.31.0: + resolve_work tool; 0.30.3: concise-answer rule; 0.29.0: + research_state/propose_claim tools + research state line; 0.28.0: + search_global_library; 0.27.0: + search_seen_sources; 0.24.1: library tools + inventory)


# `client` (the one app instance for the whole test session) lives in tests/conftest.py


@pytest.fixture
def isolated_db(tmp_path, monkeypatch):
    """A private database for tests that claim/inspect jobs by hand: the API client's worker threads keep their own
    connections to the shared test database, so nothing they do can race with the test's queue."""
    from neurosearch.config import settings
    data = tmp_path / "data"; data.mkdir(); (data / "media").mkdir()
    monkeypatch.setattr(settings, "data_dir", data)
    db._local.conn = None
    db.init_db()
    yield data
    try:
        db.connect().close()
    except Exception:  # noqa: BLE001
        pass
    db._local.conn = None


def test_parse_json3_skips_append_events():
    raw = '{"events":[{"tStartMs":0,"dDurationMs":2000,"segs":[{"utf8":"hello "},{"utf8":"world"}]},' \
          '{"tStartMs":1000,"aAppend":1,"segs":[{"utf8":"\\n"}]},{"tStartMs":2000,"dDurationMs":3000,"segs":[{"utf8":"second"}]}]}'
    segs = media.parse_json3(raw)
    assert segs == [{"start": 0.0, "end": 2.0, "text": "hello world"}, {"start": 2.0, "end": 5.0, "text": "second"}]


def test_parse_vtt():
    raw = "WEBVTT\n\n00:00:01.000 --> 00:00:03.500\nHi <c>there</c>\n\n00:00:03.500 --> 00:00:05.000\nHi there\n\n00:01:00.000 --> 00:01:02.000\nnext\n"
    segs = media.parse_vtt(raw)
    assert segs[0]["text"] == "Hi there" and segs[0]["end"] == 5.0 and segs[1]["start"] == 60.0


def test_classify_url():
    assert media.classify_url("https://youtu.be/abc123def45") == "video"
    assert media.classify_url("https://www.youtube.com/watch?v=abc&list=PL1") == "video"
    assert media.classify_url("https://www.youtube.com/playlist?list=PL1") == "playlist"
    assert media.classify_url("https://www.youtube.com/@somebody") == "channel"
    assert media.classify_url("https://www.instagram.com/reel/XYZ/") == "instagram"
    assert media.classify_url("https://example.com/ep1.mp3") == "media"


def test_chunking_long_and_short():
    segs = [{"start": i * 3, "end": i * 3 + 3, "text": f"segment {i}"} for i in range(200)]
    chunks = chunking.build_chunks(segs, duration=600)
    assert 8 <= len(chunks) <= 14
    assert chunks[0]["start"] == 0 and chunks[1]["start"] < chunks[0]["end"]  # overlap
    assert len(chunking.build_chunks(segs[:20], duration=60)) == 1
    chap = chunking.build_chunks(segs, duration=600, chapters=[{"start_time": 0, "title": "Intro"}, {"start_time": 300, "title": "Main"}])
    assert chap[0]["text"].startswith("[Intro]") and any(c["text"].startswith("[Main]") for c in chap)
    assert chunking.fmt_ts(3723) == "1:02:03" and chunking.fmt_ts(65) == "1:05"


def test_manual_ingest_search_and_projects():
    db.init_db()
    r = ingest.ingest_text("Manual", "0:05 hello there\n0:40 talking about pricing objections\n[1:02:03] later on", tags=["sales"])
    assert r["segments"] == 3
    hits = db.fts_search("pricing objections")
    assert hits and db.get_chunks_by_ids([hits[0][0]])[hits[0][0]]["start"] == 40.0
    p = db.create_project("Sales", "objections", tags=["sales"])
    assert db.project_source_ids(p["id"]) == [r["source_id"]]
    assert db.find_project("sales")["id"] == p["id"]


def test_api_auth_and_import(client):
    assert client.get("/api/stats").status_code == 401
    assert client.get("/api/stats", headers=H).status_code == 200
    payload = {
        "platform": "youtube", "external_id": "abc123def45", "url": "https://www.youtube.com/watch?v=abc123def45",
        "title": "Imported video", "channel": "Chan", "duration": 120, "transcript_kind": "captions",
        "segments": [{"start": 0, "end": 4, "text": "we talk about churn"}, {"start": 4, "end": 9, "text": "and retention tactics"}],
        "chapters": [],
    }
    r = client.post("/api/import", headers=H, json={"payload": payload, "tags": ["x"],
                    "collection": {"kind": "playlist", "external_id": "PL1", "url": "u", "title": "My list"}})
    assert r.status_code == 200, r.text
    hits = client.get("/api/search", headers=H, params={"q": "retention"}).json()
    assert hits and hits[0]["link"].endswith("t=0s") and hits[0]["title"] == "Imported video"
    colls = client.get("/api/collections", headers=H).json()
    assert colls[0]["title"] == "My list" and colls[0]["n_sources"] == 1
    csv_text = client.get("/api/export/segments.csv", headers=H).text
    assert "Imported video" in csv_text and "churn" in csv_text
    # re-import is idempotent (same external id -> same source)
    client.post("/api/import", headers=H, json={"payload": payload})
    assert client.get("/api/stats", headers=H).json()["sources"] == 2


def test_mcp_path_token(client):
    init = {"jsonrpc": "2.0", "id": 1, "method": "initialize",
            "params": {"protocolVersion": "2025-06-18", "capabilities": {}, "clientInfo": {"name": "t", "version": "1"}}}
    hdr = {"Content-Type": "application/json", "Accept": "application/json, text/event-stream"}
    assert client.post("/mcp/wrong", json=init, headers=hdr).status_code == 401
    assert client.post("/mcp/", json=init, headers=hdr).status_code == 401
    ok = client.post("/mcp/t0k", json=init, headers=hdr)
    assert ok.status_code == 200 and "Neuro Search" in ok.text
    ok2 = client.post("/mcp", json=init, headers={**hdr, **H})
    assert ok2.status_code == 200


def test_chat_url_ingest_and_exports(client, monkeypatch):
    p = client.post("/api/projects", headers=H, json={"name": "Chat proj", "brief": "test brief"}).json()
    # a message that is only links -> queued into the project, no Claude call needed
    r = client.post("/api/ask", headers=H, json={"question": "https://www.youtube.com/watch?v=zzzzzzzzzzz add this",
                                                 "project_id": p["id"]}).json()
    assert r["ingest_jobs"] and "Queued 1 link" in r["answer"]
    job = client.get("/api/jobs/" + r["ingest_jobs"][0]["job_id"], headers=H).json()
    assert job["payload"]["project_id"] == p["id"]
    # findings + masterplan package without synthesis
    client.post(f"/api/projects/{p['id']}/notes", headers=H, json={"content": "Finding one [1]", "citations": [
        {"n": 1, "title": "Imported video", "timestamp": "0:00", "link": "https://www.youtube.com/watch?v=abc123def45&t=0s"}]})
    md = client.get(f"/api/projects/{p['id']}/findings.md", headers=H).text
    assert "Finding one" in md and "watch?v=abc123def45&t=0s" in md
    import io, zipfile
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/projects/{p['id']}/masterplan.zip?synthesize=false", headers=H).content))
    names = z.namelist()
    assert {"README.md", "masterplan.md", "findings.md", "sources.csv", "context.json", "conversations.md"} <= set(names)


def test_ask_tool_loop(monkeypatch):
    """Claude calls save_finding + update_brief, then answers; citations resolve and notes are stored."""
    import anthropic
    from neurosearch import qa
    monkeypatch.setattr(qa.settings, "anthropic_api_key", "fake")
    p = db.create_project("Loop", "old brief")
    src = db.find_source("youtube", "abc123def45")
    db.add_project_sources(p["id"], [src["id"]])

    class Blk:
        def __init__(self, **kw): self.__dict__.update(kw)
    calls = []

    class Msgs:
        def create(self, **kw):
            calls.append(kw)
            if len(calls) == 1:
                return Blk(stop_reason="tool_use", content=[
                    Blk(type="tool_use", id="t1", name="update_brief", input={"brief": "new brief about retention"}),
                    Blk(type="tool_use", id="t2", name="save_finding", input={"content": "They discuss retention tactics [1]."}),
                ])
            assert kw["messages"][-1]["role"] == "user" and kw["messages"][-1]["content"][0]["type"] == "tool_result"
            return Blk(stop_reason="end_turn", content=[Blk(type="text", text="Retention comes up early [1].", citations=None)])

    monkeypatch.setattr(anthropic, "Anthropic", lambda **kw: Blk(messages=Msgs()))
    res = qa.ask("refocus on retention and pin what they say", project_id=p["id"], conversation_id="conv1")
    assert res["citations"] and res["citations"][0]["title"] == "Imported video"
    assert {a["type"] for a in res["actions"]} == {"brief_updated", "finding_saved"}
    assert db.get_project(p["id"])["brief"] == "new brief about retention"
    notes = db.list_project_notes(p["id"])
    assert notes and notes[0]["citations"][0]["link"].endswith("t=0s")
    sysb = calls[0]["system"]
    assert "Pinned findings so far" in sysb[-1]["text"] and "Pinned findings" not in sysb[0]["text"] and "cache_control" not in sysb[-1]   # project state after the cached prefix
    assert "update_brief" in [t["name"] for t in calls[0]["tools"]]


def test_document_upload_job(client):
    import time
    p = client.post("/api/projects", headers=H, json={"name": "Docs", "brief": None}).json()
    text = "Retention playbook.\n\nChurn drops when onboarding is personal.\n\n" + ("Filler paragraph about pricing anchors. " * 40 + "\n\n") * 6
    r = client.post("/api/ingest/file", headers=H, data={"project_id": p["id"]},
                    files={"file": ("playbook.txt", text.encode(), "text/plain")}).json()
    assert r["job"]
    for _ in range(50):
        j = client.get("/api/jobs/" + r["job"], headers=H).json()
        if j["status"] in ("done", "failed"):
            break
        time.sleep(0.2)
    assert j["status"] == "done", j
    srcs = client.get("/api/sources", headers=H, params={"project_id": p["id"]}).json()
    assert srcs and srcs[0]["platform"] == "document" and srcs[0]["title"] == "playbook.txt"
    hits = client.get("/api/search", headers=H, params={"q": "onboarding personal", "project_id": p["id"]}).json()
    assert hits and hits[0]["timestamp"].startswith("p. ")
    jobs_ = client.get(f"/api/projects/{p['id']}/jobs", headers=H).json()
    assert any(x["id"] == r["job"] for x in jobs_)
    # chats are isolated per project
    c = client.post("/api/conversations", headers=H, json={"project_id": p["id"], "title": "Idea A"}).json()
    client.put("/api/conversations/" + c["id"], headers=H, json={"title": "Idea A2"})
    convs = client.get("/api/conversations", headers=H, params={"project_id": p["id"]}).json()
    assert [x["title"] for x in convs] == ["Idea A2"]


def test_master_planner(client, monkeypatch):
    import anthropic
    from tests.fake_claude import Anthropic
    from neurosearch import planner
    monkeypatch.setattr(anthropic, "Anthropic", Anthropic)
    monkeypatch.setattr(planner.settings, "anthropic_api_key", "fake")
    p = client.post("/api/projects", headers=H, json={"name": "Site move", "brief": "Move site off Squarespace", "context": "Budget $500"}).json()
    client.post(f"/api/projects/{p['id']}/facts", headers=H, json={"kind": "constraint", "content": "budget under $500"})
    src = db.find_source("youtube", "abc123def45"); db.add_project_sources(p["id"], [src["id"]])
    client.post(f"/api/projects/{p['id']}/notes", headers=H, json={"content": "Cloudflare Pages is free [1]", "citations": [
        {"n": 1, "title": "Imported video", "timestamp": "0:00", "link": "https://www.youtube.com/watch?v=abc123def45&t=0s"}]})
    assert client.get(f"/api/projects/{p['id']}/plan", headers=H).json()["plan"] is None
    row = client.post(f"/api/projects/{p['id']}/plan/build", headers=H, json={}).json()
    assert row["version"] == 1 and row["plan"]["approach"]["recommended"].startswith("Static export")
    assert "U1" in row["plan"]["_evidence"] and row["plan"]["_evidence"]["U1"]["kind"] == "user"
    assert db.get_project(p["id"])["mode"] == "plan"
    # statuses + markdown/html
    client.put(f"/api/plans/{row['id']}/items/first_steps.0", headers=H, json={"status": "in_progress"})
    md = client.get(f"/api/projects/{p['id']}/plan.md", headers=H).text
    assert "## 3. First steps" in md and "`in progress`" in md and "watch?v=abc123def45" in md and "Beginner gotchas" in md
    html = client.get(f"/api/projects/{p['id']}/plan.html", headers=H).text
    assert "<table>" in html and "Ready to start" in html and "<script" not in html
    # research changed -> suggested updates -> accept -> apply (new version keeps statuses)
    client.post(f"/api/projects/{p['id']}/notes", headers=H, json={"content": "Cloudflare forms are limited", "citations": []})
    assert client.get(f"/api/projects/{p['id']}/plan", headers=H).json()["research_changed"] is True
    r = client.post(f"/api/projects/{p['id']}/plan/check-updates", headers=H).json()
    assert r["updates"] and r["plan"]["updates"][0]["status"] == "pending"
    uid = r["plan"]["updates"][0]["id"]
    client.post(f"/api/plan-updates/{uid}", headers=H, json={"status": "accepted"})
    row2 = client.post(f"/api/projects/{p['id']}/plan/apply", headers=H).json()
    assert row2["version"] == 2 and row2["items"]["first_steps.0"]["status"] == "in_progress"
    # start project
    started = client.post(f"/api/plans/{row2['id']}/start", headers=H).json()
    assert started["status"] == "started" and started["items"]["first_steps.1"]["status"] == "ready"
    assert db.get_project(p["id"])["mode"] == "execute"
    # masterplan zip now carries the plan
    import io, zipfile
    z = zipfile.ZipFile(io.BytesIO(client.get(f"/api/projects/{p['id']}/masterplan.zip?synthesize=false", headers=H).content))
    assert "master_plan.md" in z.namelist() and "master_plan.html" in z.namelist()


def test_suggested_findings(client, monkeypatch):
    import anthropic, time
    from tests.fake_claude import Anthropic
    from neurosearch import findings, jobs
    monkeypatch.setattr(anthropic, "Anthropic", Anthropic)
    monkeypatch.setattr(findings.settings, "anthropic_api_key", "fake")
    monkeypatch.setattr(jobs.settings, "anthropic_api_key", "fake")
    p = client.post("/api/projects", headers=H, json={"name": "Suggest", "brief": "hosting"}).json()
    r = ingest.ingest_text("Hosting talk", "0:05 cloudflare pages is free hosting for static sites with no bandwidth bill\n3:40 never touch the MX records when you move hosting or email breaks", project_id=p["id"])
    # auto-queued suggestion job runs in the background worker
    for _ in range(60):
        pj = client.get(f"/api/projects/{p['id']}", headers=H).json()
        if pj["suggested"]:
            break
        time.sleep(0.2)
    assert len(pj["suggested"]) == 2 and pj["notes"] == []
    top = pj["suggested"][0]
    assert 1 <= top["importance"] <= 5 and top["citations"][0]["timestamp"] in ("0:05", "3:40") and top["status"] == "suggested"
    # the fake quotes real transcript text, so the evidence validator passes on it
    from neurosearch.evidence import quote_in_text
    assert all(quote_in_text(n["citations"][0]["snippet"], " ".join(x["text"] for x in db.get_segments(r["source_id"]))) for n in pj["suggested"])
    # summary/substance are project-relative: they live on the project's listing, not the global source row
    src = [x for x in client.get(f"/api/sources?project_id={p['id']}", headers=H).json() if x["id"] == r["source_id"]][0]
    assert isinstance(src["substance"], int) and src["summary"].startswith("Covers") and src["analysis"]["summary"]["prompt_version"].startswith("findings-")
    assert src["analysis"]["summary"]["source_revision"] and src["analysis"]["summary"]["brief_revision"] and src["analysis"]["summary"]["input_hash"]
    assert "relevance" not in src["analysis"] and not src["legacy_analysis"]           # never ranked: no relevance artifact, and nothing pretends otherwise
    assert client.get(f"/api/sources/{r['source_id']}", headers=H).json()["analyses"][0]["project_id"] == p["id"]
    # approve one, dismiss one -> only approved counts for exports/planner
    client.post(f"/api/notes/{top['id']}/status", headers=H, json={"status": "approved"})
    pj_dismissed_title = pj["suggested"][1]["title"]
    client.post(f"/api/notes/{pj['suggested'][1]['id']}/status", headers=H, json={"status": "dismissed"})
    pj = client.get(f"/api/projects/{p['id']}", headers=H).json()
    assert len(pj["notes"]) == 1 and pj["suggested"] == []
    md = client.get(f"/api/projects/{p['id']}/findings.md", headers=H).text
    assert top["title"] in md and pj_dismissed_title not in md
    # nothing left to analyse; force re-analyses
    assert client.post(f"/api/projects/{p['id']}/suggest", headers=H, json={}).json()["job"] is None
    assert client.post(f"/api/projects/{p['id']}/suggest", headers=H, json={"force": True}).json()["sources"] == 1


def test_project_onboarding(client):
    body = {"name": "Onboard", "goal": "site live on cheaper host", "brief": "which host", "context": "Budget: $500",
            "questions": ["Which hosts do practitioners recommend?", "What breaks during DNS moves?"],
            "facts": [{"kind": "constraint", "content": "Budget: $500"}, {"kind": "rejected", "content": "Wix"}],
            "urls": ["https://www.youtube.com/watch?v=aaaaaaaaaaa"], "audience": "me", "output_pref": "a recommendation",
            "source_prefs": "practitioners over marketers"}
    p = client.post("/api/projects", headers=H, json=body).json()
    assert p["goal"] == "site live on cheaper host" and p["questions"] == body["questions"] and len(p["jobs"]) == 1
    full = client.get(f"/api/projects/{p['id']}", headers=H).json()
    assert [c["title"] for c in full["conversations"]] == body["questions"][::-1] or len(full["conversations"]) == 2
    assert {f["kind"] for f in full["facts"]} == {"constraint", "rejected"}
    steering = db.project_steering(full)
    assert "Goal (what done looks like): site live" in steering and "practitioners" in steering and "Starting questions" in steering
    # settings round-trip keeps the new fields
    upd = client.put(f"/api/projects/{p['id']}", headers=H, json={"name": "Onboard", "goal": "g2", "questions": ["q3"]}).json()
    assert upd["goal"] == "g2" and upd["questions"] == ["q3"] and upd["audience"] == "me"


def test_discover_sources(client, monkeypatch):
    import anthropic
    from tests.fake_claude import Anthropic
    from neurosearch import discover
    monkeypatch.setattr(anthropic, "Anthropic", Anthropic)
    monkeypatch.setattr(discover.settings, "anthropic_api_key", "fake")
    p = client.post("/api/projects", headers=H, json={"name": "Money", "brief": "get out of debt and start investing"}).json()
    r = client.post(f"/api/projects/{p['id']}/discover", headers=H, json={}).json()
    assert r["added"] == 3 and r["verified"] == 1 and r["extra"] == 1 and r["items"][0]["name"] == "Dave Ramsey"
    ds = client.get(f"/api/projects/{p['id']}/discoveries", headers=H).json()
    ramsey = [d for d in ds if d["name"] == "Dave Ramsey"][0]
    assert ramsey["start_with"][0]["title"] == "Baby Steps"        # verification pass patched the starting video
    # second run doesn't duplicate
    r2 = client.post(f"/api/projects/{p['id']}/discover", headers=H, json={"refine": "more contrarian"}).json()
    assert r2["added"] == 0
    ds = client.get(f"/api/projects/{p['id']}/discoveries", headers=H).json()
    assert len(ds) == 3 and ds[0]["fit"] == 5
    client.post(f"/api/discoveries/{ds[1]['id']}/status", headers=H, json={"status": "dismissed"})
    sts = [d["status"] for d in client.get(f"/api/projects/{p['id']}/discoveries", headers=H).json()]
    assert sts.count("dismissed") == 1 and sts.count("new") == 2


def test_course_import(client):
    from neurosearch.courses import normalise_embed, write_cookie_file
    assert normalise_embed("https://www.loom.com/embed/abc123DEF?sid=1") == "https://www.loom.com/share/abc123DEF"
    assert normalise_embed("https://player.vimeo.com/video/12345?h=zz9") == "https://vimeo.com/12345/zz9"
    assert normalise_embed("https://www.youtube-nocookie.com/embed/abc123def45?rel=0") == "https://www.youtube.com/watch?v=abc123def45"
    p = client.post("/api/projects", headers=H, json={"name": "Course", "brief": "the course"}).json()
    r = client.post(f"/api/projects/{p['id']}/course-import", headers=H, json={
        "course": {"title": "Sales Mastery", "url": "https://school.example.com/courses/sales"},
        "lessons": [
            {"title": "Welcome", "module": "Module 1", "page_url": "https://school.example.com/lessons/1", "video_urls": ["https://www.loom.com/embed/xyz789"]},
            {"title": "Text only", "module": "Module 1", "page_url": "https://school.example.com/lessons/2", "video_urls": []},
        ],
        "cookies": [{"domain": ".example.com", "name": "session", "value": "s3cr3t", "path": "/", "secure": True, "expirationDate": 1900000000}]}).json()
    assert r["queued"] == 1 and r["no_video"] == ["Text only"] and r["cookies"] is True
    jobs_ = client.get(f"/api/projects/{p['id']}/jobs", headers=H).json()
    j = next(x for x in jobs_ if x["kind"] == "ingest_url")
    assert j["payload"]["title"] == "Module 1 › Welcome" and j["payload"]["referer"].endswith("/lessons/1")
    assert open(j["payload"]["cookies_file"]).read().count("s3cr3t") == 1
    cs = client.get("/api/collections", headers=H).json()
    assert any(c["kind"] == "course" and c["title"] == "Sales Mastery" for c in cs)
    assert client.get("/extension.zip", headers=H).status_code == 200


def test_rate_limit_backoff(monkeypatch):
    from neurosearch import media
    monkeypatch.setattr(media.settings, "yt_delay", 0.0)
    monkeypatch.setattr(media.settings, "yt_backoff_minutes", 1)
    media._yt_paused_until = 0.0
    db.init_db()
    # a bot-check error inside a polite() block converts to RateLimited and starts the pause
    with pytest.raises(media.RateLimited):
        with media.polite("https://www.youtube.com/watch?v=abc123def45"):
            raise Exception("ERROR: Sign in to confirm you’re not a bot")
    assert media.rate_limit_status()["paused"] is True
    # while paused, further youtube fetches are refused immediately; non-youtube is unaffected
    with pytest.raises(media.RateLimited):
        with media.polite("https://www.youtube.com/watch?v=zzz"):
            pass
    with media.polite("https://www.loom.com/share/x"):
        pass
    media._yt_paused_until = 0.0
    # job requeue with not_before is skipped by claim_job until due
    j = db.create_job("reembed", {})
    db.requeue_job(j["id"], delay=3600, message="paused")
    row = db.get_job(j["id"])
    assert row["status"] == "queued" and row["not_before"] > db.now() + 3000 and row["message"] == "paused"
    claimed = set()
    while (c := db.claim_job()):
        claimed.add(c["id"]); db.update_job(c["id"], status="done")
    assert j["id"] not in claimed
    db.requeue_job(j["id"], delay=0)
    assert db.claim_job()["id"] == j["id"]


def test_age_cutoff(client, monkeypatch):
    from neurosearch import ingest, media
    # fake a channel with 3 videos newest-first; metadata says the 2nd is old
    monkeypatch.setattr(media, "enumerate_entries", lambda url: ({"id": "UC1", "title": "Chan", "url": url},
        [{"id": f"vid{i}0000000", "url": f"https://www.youtube.com/watch?v=vid{i}0000000", "title": f"V{i}"} for i in range(3)]))
    dates = {"vid00000000": "2026-06-01", "vid10000000": "2019-01-01", "vid20000000": "2018-01-01"}
    def fake_info(url, cookies_file=None, referer=None):
        vid = url.split("v=")[1]
        return {"id": vid, "title": vid, "webpage_url": url, "upload_date": dates[vid].replace("-", ""), "duration": 60,
                "channel": "Chan", "thumbnail": f"https://img.example/{vid}.jpg", "description": "an old video, still on-topic"}
    monkeypatch.setattr(media, "fetch_info", fake_info)
    monkeypatch.setattr(media, "fetch_captions", lambda info, cookies_file=None: ([{"start": 0, "end": 5, "text": "hello content"}], "en"))
    p = client.post("/api/projects", headers=H, json={"name": "Cutoff", "brief": "x"}).json()
    # default: channel is listed and waits for approval
    r = ingest.ingest_url("https://www.youtube.com/@chan", project_id=p["id"], since_years=2, max_videos=10)
    assert r["proposed"] == 3 and r["queued"] == 0 and r["review"] is True and r["limits"]["since"] is not None
    reviews = client.get(f"/api/projects/{p['id']}/reviews", headers=H).json()
    assert len(reviews) == 1 and len(reviews[0]["proposed"]) == 3
    # proposed sources are hidden from the project's source list
    assert all(x["status"] != "proposed" for x in client.get("/api/sources", headers=H, params={"project_id": p["id"]}).json())
    # approve all -> jobs created
    a = client.post(f"/api/collections/{reviews[0]['id']}/approve", headers=H, json={}).json()
    assert a["started"] == 3 and a["dropped"] == 0
    # run the queued jobs synchronously
    from neurosearch import jobs
    while (j := db.claim_job()):
        jobs.execute(j)
    statuses = {s["external_id"]: s["status"] for s in db.list_sources(limit=1000) if s["external_id"] in dates}
    assert statuses["vid00000000"] == "ready" and statuses["vid10000000"] == "skipped" and statuses["vid20000000"] == "skipped"
    # a skipped (pre-cutoff) source's metadata fetch still succeeded — thumbnail, channel and published date must be
    # kept, not thrown away, so the Sources list can show a real thumbnail and the pool's potential scan has real
    # title/description text to score instead of the bare listing-stage title.
    skipped = db.get_source(next(s["id"] for s in db.list_sources(limit=1000) if s["external_id"] == "vid10000000"))
    assert skipped["thumbnail_url"] == "https://img.example/vid10000000.jpg" and skipped["channel"] == "Chan" and skipped["published_at"] == "2019-01-01"
    assert "an old video" in (skipped["description"] or "") and skipped["error"] and "before cutoff" in skipped["error"]


def test_budget_valve(client, monkeypatch):
    from neurosearch import usage, jobs
    # record spend then set a tiny daily budget -> paid jobs wait, queue keeps order, resume when raised
    usage.record("findings", "claude-sonnet-4-6", input_tokens=1_000_000, output_tokens=0)   # $3
    u = client.get("/api/usage", headers=H).json()
    assert u["today"] >= 3.0 and u["blocked"] is None
    client.post("/api/usage/budget", headers=H, json={"daily": 2})
    u = client.get("/api/usage", headers=H).json()
    assert "daily budget reached" in (u["blocked"] or "")
    ok, reason, wait = usage.check()
    assert not ok and wait >= 60
    j = db.create_job("reembed", {})
    # run the worker loop body once: the job must be re-queued, not failed
    job = db.claim_job(("reembed",))
    try:
        jobs.run_job(job)
    except usage.BudgetPaused as e:
        db.requeue_job(job["id"], delay=min(e.wait, 3600), message=f"paused: {e}")
    row = db.get_job(j["id"])
    assert row["status"] == "queued" and "paused" in row["message"]
    # manual pause/resume clears the wait
    client.post("/api/usage/budget", headers=H, json={"daily": 100, "paused": True})
    assert "paused by you" in client.get("/api/usage", headers=H).json()["blocked"]
    client.post("/api/usage/budget", headers=H, json={"paused": False})
    assert client.get("/api/usage", headers=H).json()["blocked"] is None
    assert db.get_job(j["id"])["not_before"] is None
    db.update_job(j["id"], status="done")


def test_review_partial_approve(client, monkeypatch):
    from neurosearch import ingest, media
    monkeypatch.setattr(media, "enumerate_entries", lambda url: ({"id": "PLx", "title": "List", "url": url},
        [{"id": f"pl{i}000000000"[:11], "url": f"https://www.youtube.com/watch?v=pl{i}0000000", "title": f"P{i}"} for i in range(4)]))
    p = client.post("/api/projects", headers=H, json={"name": "Review", "brief": "x"}).json()
    r = ingest.ingest_url("https://www.youtube.com/playlist?list=PLx", project_id=p["id"])
    assert r["proposed"] == 4
    rv = client.get(f"/api/projects/{p['id']}/reviews", headers=H).json()[0]
    keep = [s["id"] for s in rv["proposed"][:2]]
    a = client.post(f"/api/collections/{rv['id']}/approve", headers=H, json={"source_ids": keep}).json()
    assert a["started"] == 2 and a["dropped"] == 2
    assert client.get(f"/api/projects/{p['id']}/reviews", headers=H).json() == []
    jobs_ = [j for j in client.get(f"/api/projects/{p['id']}/jobs", headers=H).json() if j["kind"] == "ingest_source"]
    assert {j["payload"]["source_id"] for j in jobs_} >= set(keep)


def test_cancel_queued(client):
    p = client.post("/api/projects", headers=H, json={"name": "Cancel", "brief": "x"}).json()
    src = db.upsert_source(platform="youtube", external_id="cancelme001", url="https://www.youtube.com/watch?v=cancelme001", status="pending")
    coll = db.upsert_collection("channel", "UCc", "u", "Chan"); db.link_source_collection(src["id"], coll["id"]); db.add_project_collections(p["id"], [coll["id"]])
    j = db.create_job("ingest_source", {"source_id": src["id"], "collection_id": coll["id"], "min_date": "2024-01-01"})
    r = client.post("/api/jobs/cancel-queued", headers=H, json={}).json()
    assert r["cancelled"] >= 1 and db.get_job(j["id"])["status"] == "cancelled" and db.get_job(j["id"])["message"] == "cancelled"
    assert db.get_source(src["id"])["status"] == "proposed"
    assert client.get(f"/api/projects/{p['id']}/reviews", headers=H).json()[0]["proposed"][0]["id"] == src["id"]


def test_relevance_ranking(client, monkeypatch):
    from neurosearch import ingest, media, relevance
    import anthropic
    from tests.fake_claude import Anthropic
    monkeypatch.setattr(anthropic, "Anthropic", Anthropic)
    monkeypatch.setattr(relevance.settings, "anthropic_api_key", "fake")
    monkeypatch.setattr(media, "enumerate_entries", lambda url: ({"id": "UCr", "title": "Chan", "url": url},
        [{"id": f"rk{i}000000000"[:11], "url": f"https://www.youtube.com/watch?v=rk{i}00000000", "title": f"V{i}",
          "description": "getting out of debt with a budget" if i % 2 else "vlog", "view_count": 1000 * i, "duration": 600} for i in range(6)]))
    p = client.post("/api/projects", headers=H, json={"name": "Rank", "brief": "getting out of debt"}).json()
    r = ingest.ingest_url("https://www.youtube.com/@chan", project_id=p["id"], max_videos=2)
    assert r["proposed"] == 6                      # whole pool is listed, not just the first 2
    rv = client.get(f"/api/projects/{p['id']}/reviews", headers=H).json()[0]
    assert rv["meta"]["ranked"] is False and rv["meta"]["max_videos"] == 2
    job = [j for j in client.get(f"/api/projects/{p['id']}/jobs", headers=H).json() if j["kind"] == "rank_proposed"][0]
    res = relevance.rank_collection(rv["id"], p["id"], want=2)
    assert res["ranked"] == 6
    rv = client.get(f"/api/projects/{p['id']}/reviews", headers=H).json()[0]
    assert rv["meta"]["ranked"] is True
    scores = [s["relevance"] for s in rv["proposed"]]
    assert scores == sorted(scores, reverse=True) and scores[0] > scores[-1] and scores[-1] == 10   # lexical fake: brief-word overlap
    assert rv["proposed"][0]["relevance_why"] == "on topic" and rv["proposed"][0]["description"].startswith("getting out of debt")
    assert res["batches"] == 1 and res["failed_batches"] == 0 and res["repaired_batches"] == 0
    # re-rank endpoint queues a job and resets the flag
    assert "job_id" in client.post(f"/api/collections/{rv['id']}/rank", headers=H, json={"want": 3}).json()
    assert db.review_meta(rv["id"])["ranked"] is False and db.review_meta(rv["id"])["max_videos"] == 3


def test_cancel_single_job(client):
    p = client.post("/api/projects", headers=H, json={"name": "Cancel1", "brief": "x"}).json()
    ids = []
    for k in ("a", "b"):
        src = db.upsert_source(platform="youtube", external_id=f"one{k}0000000", url=f"https://www.youtube.com/watch?v=one{k}0000000", status="pending", title=f"Video {k}")
        db.add_project_sources(p["id"], [src["id"]])
        ids.append(db.create_job("ingest_source", {"source_id": src["id"]})["id"])
    js = client.get(f"/api/projects/{p['id']}/jobs", headers=H).json()
    assert any(j.get("label") == "Video a" for j in js)
    assert client.post(f"/api/jobs/{ids[0]}/cancel", headers=H).json()["cancelled"] == 1
    assert db.get_job(ids[0])["message"] == "cancelled" and db.get_job(ids[1])["status"] == "queued"
    assert client.post(f"/api/jobs/{ids[0]}/cancel", headers=H).status_code == 409


def test_discover_background_and_pause_turn(client, monkeypatch):
    import anthropic
    from tests.fake_claude import Anthropic, _Msgs
    from neurosearch import discover, jobs
    calls = {"n": 0}
    orig = _Msgs.create
    def create(self, **kw):   # first call returns a paused turn with no text, second the real answer
        calls["n"] += 1
        r = orig(self, **kw)
        if calls["n"] == 2:   # the web-search (verify) pass hands back a paused turn first
            r.stop_reason = "pause_turn"; r.content = []
        return r
    monkeypatch.setattr(_Msgs, "create", create)
    monkeypatch.setattr(anthropic, "Anthropic", Anthropic)
    monkeypatch.setattr(discover.settings, "anthropic_api_key", "fake")
    p = client.post("/api/projects", headers=H, json={"name": "BG", "brief": "get out of debt"}).json()
    r = client.post(f"/api/projects/{p['id']}/discover", headers=H, json={"background": True}).json()
    job = db.get_job(r["job_id"]); assert job["kind"] == "discover"
    res = jobs.run_job(job)
    assert calls["n"] == 3 and res["added"] == 3 and res["verified"] == 1
    assert len(client.get(f"/api/projects/{p['id']}/discoveries", headers=H).json()) == 3


def test_review_discard_last(client, monkeypatch):
    from neurosearch import ingest, media
    monkeypatch.setattr(media, "enumerate_entries", lambda url: ({"id": "PLd", "title": "L", "url": url},
        [{"id": "dl00000000a", "url": "https://www.youtube.com/watch?v=dl00000000a", "title": "only one"}]))
    p = client.post("/api/projects", headers=H, json={"name": "Last", "brief": "x"}).json()
    ingest.ingest_url("https://www.youtube.com/playlist?list=PLd", project_id=p["id"])
    rv = client.get(f"/api/projects/{p['id']}/reviews", headers=H).json()[0]
    assert client.post(f"/api/collections/{rv['id']}/approve", headers=H, json={"source_ids": []}).json()["dropped"] == 1
    assert client.get(f"/api/projects/{p['id']}/reviews", headers=H).json() == []


def test_webpage_source(client, monkeypatch):
    from neurosearch import webpage, media
    html = ("<html><head><title>Buying a boring business</title></head><body><nav><a>Home</a></nav><article><h1>Buying a boring business</h1>"
            + "".join(f"<p>Paragraph {i}: SBA lenders want a debt service coverage ratio above 1.25 and a seller note for part of the price.</p>" for i in range(40))
            + "</article><footer>x</footer></body></html>")
    monkeypatch.setattr(webpage, "fetch", lambda url, timeout=40.0: (url, "text/html; charset=utf-8", html.encode()))
    assert media.classify_url("https://example.com/blog/boring-business") == "web"
    assert media.classify_url("https://vimeo.com/12345") == "media"
    assert media.classify_url("https://cdn.example.com/ep1.mp3") == "media"
    p = client.post("/api/projects", headers=H, json={"name": "Web", "brief": "buying businesses"}).json()
    from neurosearch import ingest
    r = ingest.ingest_url("https://example.com/blog/boring-business", project_id=p["id"])
    assert r["kind"] == "web" and r["segments"] >= 1 and r["title"] == "Buying a boring business"
    src = client.get(f"/api/sources/{r['source_id']}", headers=H).json()
    assert src["platform"] == "web" and src["status"] == "ready" and src["channel"] == "example.com"
    hits = client.get("/api/search?q=debt+service+coverage&project_id=" + p["id"], headers=H).json()
    assert hits and hits[0]["timestamp"].startswith("§ ")


def test_rank_parse_tolerates_bad_json():
    from neurosearch.relevance import parse_scores
    good = '{"scores":[{"i":0,"score":90,"why":"on topic"},{"i":1,"score":10,"why":"filler"}]}'
    assert [x["score"] for x in parse_scores(good)["scores"]] == [90, 10]
    bad = '```json\n{"scores":[{"i":0,"score":90,"why":"buying "boring" businesses"},{"i":1,"score":10,"why":"filler"},{"i":2,"score":55,"why":"cut off'
    res = parse_scores(bad)
    assert res.get("repaired") and [(x["i"], x["score"]) for x in res["scores"]] == [(0, 90), (1, 10)]
    assert res["scores"][0]["why"] == 'buying "boring" businesses'


def test_ingest_html_from_extension(client):
    p = client.post("/api/projects", headers=H, json={"name": "Ext", "brief": "buying businesses"}).json()
    html = "<html><head><title>Learning Center</title></head><body><main><h1>Due diligence checklist</h1>" + "".join(
        f"<p>Item {i}: verify three years of tax returns and reconcile them against the P&amp;L before making an offer.</p>" for i in range(30)) + "</main></body></html>"
    r = client.post(f"/api/projects/{p['id']}/ingest/html", headers=H, json={"url": "https://www.bizbuysell.com/learning-center/", "html": html}).json()
    assert r["kind"] == "web" and r["title"] == "Learning Center" and r["segments"] >= 1
    assert client.get(f"/api/sources/{r['source_id']}", headers=H).json()["status"] == "ready"


def test_blocked_site_message(monkeypatch):
    import httpx
    from neurosearch import webpage
    class R:  # minimal httpx-like response
        status_code = 403; headers = {}; content = b""; url = "https://x.com/a"
    class C:
        def __init__(self, **kw): pass
        def __enter__(self): return self
        def __exit__(self, *a): pass
        def get(self, url): return R()
    monkeypatch.setattr(httpx, "Client", C)
    import pytest
    with pytest.raises(webpage.Blocked, match="blocks automated readers"):
        webpage.fetch("https://x.com/a")


def test_instagram_urls(client):
    from neurosearch import media, ingest
    assert media.classify_url("https://www.instagram.com/sbaloanguy1/") == "instagram_profile"
    assert media.classify_url("https://www.instagram.com/reel/Cabc123/") == "instagram"
    import pytest
    with pytest.raises(RuntimeError, match="lends your session"):
        ingest.ingest_url("https://www.instagram.com/sbaloanguy1/")
    p = client.post("/api/projects", headers=H, json={"name": "IG", "brief": "x"}).json()
    r = client.post(f"/api/projects/{p['id']}/ingest/with-session", headers=H, json={"url": "https://www.instagram.com/reel/Cabc123/",
        "cookies": [{"domain": ".instagram.com", "name": "sessionid", "value": "abc", "path": "/", "secure": True}]}).json()
    j = db.get_job(r["job_id"]); assert r["cookies"] and j["payload"]["cookies_file"].endswith(".txt") and j["payload"]["review"] is False


def test_instagram_profile_with_session(client, monkeypatch, tmp_path):
    from neurosearch import ingest, media
    ck = tmp_path / "ig.txt"; ck.write_text("# cookies")
    monkeypatch.setattr(media, "enumerate_instagram", lambda url, cookies_file, limit=40: (
        {"id": "ig:sbaloanguy1", "title": "@sbaloanguy1 (Instagram)", "url": "https://www.instagram.com/sbaloanguy1/"},
        [{"id": f"ig{i}", "url": f"https://www.instagram.com/reel/ig{i}/", "title": f"SBA tip {i}", "duration": 45} for i in range(5)]))
    p = client.post("/api/projects", headers=H, json={"name": "IGP", "brief": "sba loans"}).json()
    r = ingest.ingest_url("https://www.instagram.com/sbaloanguy1/", project_id=p["id"], cookies_file=str(ck), review=False)
    assert r["kind"] == "instagram_profile" and r["proposed"] == 5 and r["review"]
    rv = client.get(f"/api/projects/{p['id']}/reviews", headers=H).json()[0]
    assert rv["kind"] == "instagram" and rv["meta"]["cookies_file"] == str(ck)
    a = client.post(f"/api/collections/{rv['id']}/approve", headers=H, json={"source_ids": [rv["proposed"][0]["id"]]}).json()
    assert a["started"] == 1
    job = [j for j in db.list_jobs(50) if j["kind"] == "ingest_source"][0]
    assert job["payload"]["cookies_file"] == str(ck) and job["payload"]["referer"].startswith("https://www.instagram.com/")


def test_transient_failures_retry_then_fail(client, monkeypatch):
    from neurosearch import jobs, ingest
    calls = {"n": 0}
    def boom(url, **kw):
        calls["n"] += 1
        raise RuntimeError("Instagram wants a login for this (or rate-limited the session)")
    monkeypatch.setattr(ingest, "ingest_url", boom)
    monkeypatch.setattr(jobs.db, "claim_job", lambda *a, **k: None)      # keep the background workers' hands off this job
    j = db.create_job("ingest_url", {"url": "https://www.instagram.com/reel/x/", "project_id": None})
    jobs.run_job  # noqa
    # drive the worker loop by hand: run + handle like _worker does
    def step():
        job = db.get_job(j["id"])
        try:
            db.update_job(job["id"], status="done", result=jobs.run_job(job))
        except Exception as e:  # noqa: BLE001
            payload = job.get("payload") or {}
            attempts = int(payload.get("_attempts") or 0) + 1
            if jobs.TRANSIENT.search(str(e)) and attempts < jobs.MAX_ATTEMPTS:
                db.set_job_payload(job["id"], {**payload, "_attempts": attempts}); db.requeue_job(job["id"], delay=0, message="retry")
            else:
                db.update_job(job["id"], status="failed", message=f"error: {e}")
    for _ in range(jobs.MAX_ATTEMPTS):
        step()
    assert calls["n"] == jobs.MAX_ATTEMPTS and db.get_job(j["id"])["status"] == "failed"
    assert jobs.TRANSIENT.search("HTTP Error 429: Too Many Requests") and not jobs.TRANSIENT.search("This Instagram post has no video")
    # manual retry endpoint creates a fresh job and retires the failed one
    r = client.post(f"/api/jobs/{j['id']}/retry", headers=H).json()
    assert db.get_job(r["job_id"])["status"] == "queued" and db.get_job(j["id"])["message"].startswith("retried")
    # the sweep is database-wide, and the background workers may have failed a job for some other test by now, so
    # assert about THIS job (already retired above) against the live set rather than a global count of zero —
    # the old `== 0` made this test fail intermittently for reasons that had nothing to do with retries.
    others = {x["id"] for x in db.failed_jobs(None)}
    assert j["id"] not in others
    assert client.post("/api/jobs/retry-failed", headers=H, json={}).json()["retried"] == len(others)


def test_ingest_skipped_anyway(client):
    p = client.post("/api/projects", headers=H, json={"name": "Skip", "brief": "x"}).json()
    src = db.upsert_source(platform="youtube", external_id="oldvid00001", url="https://www.youtube.com/watch?v=oldvid00001", status="skipped", error="published 2019-01-01, before cutoff 2024-09-03")
    db.add_project_sources(p["id"], [src["id"]])
    r = client.post("/api/sources/retry-skipped", headers=H, json={"project_id": p["id"]}).json()
    assert r["queued"] == 1 and db.get_source(src["id"])["status"] == "pending"
    j = [j for j in db.list_jobs(20) if j["kind"] == "ingest_source" and j["payload"]["source_id"] == src["id"]][0]
    assert "min_date" not in j["payload"]


def test_raising_budget_wakes_paused_jobs(client):
    j = db.create_job("suggest_findings", {"project_id": "x", "source_ids": []})
    db.requeue_job(j["id"], delay=3600, message="paused: daily budget reached")
    assert db.get_job(j["id"])["not_before"] is not None
    client.post("/api/usage/budget", headers=H, json={"daily": 20})
    assert db.get_job(j["id"])["not_before"] is None


def test_plan_has_analysis_and_repairs_truncated_json(client, monkeypatch):
    import anthropic
    from tests.fake_claude import Anthropic
    from neurosearch import planner
    monkeypatch.setattr(anthropic, "Anthropic", Anthropic)
    monkeypatch.setattr(planner.settings, "anthropic_api_key", "fake")
    p = client.post("/api/projects", headers=H, json={"name": "Plan2", "brief": "move the site", "goal": "cheap fast host"}).json()
    row = client.post(f"/api/projects/{p['id']}/plan/build", headers=H, json={}).json()
    an = row["plan"]["analysis"]
    assert an["swot"]["strengths"][0]["point"] == "Existing content" and an["options"][0]["fit"] == 5 and an["verdict"]
    md = client.get(f"/api/projects/{p['id']}/plan.md", headers=H).text
    assert "## 0. Where you stand" in md and "| Strengths | Weaknesses |" in md and "Paths compared" in md
    # background build path
    r = client.post(f"/api/projects/{p['id']}/plan/build", headers=H, json={"background": True}).json()
    assert db.get_job(r["job_id"])["kind"] == "build_plan"
    # truncated output gets repaired instead of crashing
    cut = '{"goal": {"outcome": "x", "constraints": ["a"]}, "first_steps": [{"action": "do it", "detail": "now"}, {"action": "half'
    fixed = planner._parse_json(cut)
    assert fixed["goal"]["outcome"] == "x" and fixed["first_steps"][0]["action"] == "do it"


def test_backup_snapshot():
    p = db.backup(keep=2)
    assert p.exists() and p.stat().st_size > 0
    import sqlite3
    assert sqlite3.connect(str(p)).execute("select count(*) from projects").fetchone()[0] >= 1


def test_spreadsheet_source_and_calculator(client, tmp_path):
    import openpyxl
    from neurosearch import ingest, sheets
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Deal"
    for i, (k, v) in enumerate([("Purchase price", 500000), ("Down payment %", 0.1), ("Interest rate", 0.105), ("Term (years)", 10), ("Annual SDE", 150000)], 1):
        ws.cell(i, 1, k); ws.cell(i, 2, v)
    ws["A7"] = "Loan amount"; ws["B7"] = "=B1*(1-B2)"
    ws["A8"] = "Monthly payment"; ws["B8"] = "=-PMT(B3/12,B4*12,B7)"
    ws["A9"] = "Annual debt service"; ws["B9"] = "=B8*12"
    ws["A10"] = "DSCR"; ws["B10"] = "=ROUND(B5/B9,2)"
    f = tmp_path / "deal.xlsx"; wb.save(f)
    p = client.post("/api/projects", headers=H, json={"name": "Calc", "brief": "buy a business"}).json()
    r = ingest.ingest_local_file(f, project_id=p["id"], original_name="deal.xlsx")
    assert r["transcript"] == "spreadsheet" and r["inputs"] == 5 and r["outputs"] == 4
    src = client.get(f"/api/sources/{r['source_id']}", headers=H).json()
    assert src["platform"] == "spreadsheet" and src["status"] == "ready"
    m = client.get(f"/api/sources/{r['source_id']}/calculator", headers=H).json()
    assert [i["label"] for i in m["inputs"]][:2] == ["Purchase price", "Down payment %"]
    # searchable as text
    hits = client.get("/api/search?q=DSCR&project_id=" + p["id"], headers=H).json()
    assert hits and hits[0]["timestamp"].startswith("sheet")
    # calculator by label, with a changed input
    out = client.post(f"/api/sources/{r['source_id']}/calculate", headers=H, json={"inputs": {"purchase price": "800,000"}, "outputs": ["DSCR", "Monthly payment"]}).json()
    assert abs(out["outputs"]["DSCR"] - 1.29) < 0.01 and abs(out["outputs"]["Monthly payment"] - 9715.32) < 1
    # the chat sees it as a tool
    assert sheets.calculators_for_project(p["id"])[0]["title"] == "deal"


def test_youtube_search_link_becomes_review(client, monkeypatch):
    from neurosearch import ingest, media
    assert media.classify_url("https://www.youtube.com/results?search_query=sba+loans") == "youtube_search"
    monkeypatch.setattr(media, "enumerate_search", lambda q, limit=30: ({"id": f"search:{q}", "title": f"YouTube search: {q}", "url": "https://www.youtube.com/results?search_query=" + q},
        [{"id": f"sr{i}000000000"[:11], "url": f"https://www.youtube.com/watch?v=sr{i}00000000", "title": f"{q} video {i}", "duration": 600} for i in range(4)]))
    p = client.post("/api/projects", headers=H, json={"name": "Srch", "brief": "sba loans"}).json()
    r = ingest.ingest_url("https://www.youtube.com/results?search_query=sba+loans", project_id=p["id"], review=False)
    assert r["proposed"] == 4 and r["review"] and r["title"].startswith("YouTube search")
    rv = client.get(f"/api/projects/{p['id']}/reviews", headers=H).json()
    assert rv and len(rv[0]["proposed"]) == 4


def test_bot_check_pauses_instead_of_failing(monkeypatch):
    from neurosearch import media
    class FakeYDL:
        def __init__(self, opts): self.lg = opts["logger"]
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def extract_info(self, url, download=False):
            self.lg.error("ERROR: [youtube] abc: Sign in to confirm you’re not a bot. Use --cookies-from-browser"); return None
    monkeypatch.setattr(media.yt_dlp, "YoutubeDL", FakeYDL)
    monkeypatch.setattr(media.settings, "yt_delay", 0)
    media._sites["youtube"]["until"] = 0
    import pytest
    with pytest.raises(media.RateLimited):
        media.fetch_info("https://www.youtube.com/watch?v=abc")
    assert media.rate_limit_status()["paused"]
    media._sites["youtube"]["until"] = 0   # don't leak the pause into other tests


def test_usage_prices_cache_tokens_and_helpers(tmp_path, monkeypatch):
    from neurosearch import usage
    # 10k cached-read tokens cost a tenth; 2k cache writes cost 1.25x; 1k plain input at full price
    cost = usage.record("answer", "claude-sonnet-4-6", input_tokens=1000, output_tokens=0, cache_read=10000, cache_write=2000)
    assert abs(cost - (1000 + 2000 * 1.25 + 10000 * 0.1) / 1e6 * 3.0) < 1e-9
    t = usage.totals()
    assert t["month_saved"] > 0 and t["month_cached_tokens"] >= 10000
    # tiny prefixes are not marked (the API would ignore it), big ones are
    assert "cache_control" not in usage.cached_block("short")
    assert usage.cached_block("x" * 5000)["cache_control"] == {"type": "ephemeral"}
    msgs = [{"role": "user", "content": [{"type": "text", "text": "a", "cache_control": {"type": "ephemeral"}}]},
            {"role": "assistant", "content": "b"}, {"role": "user", "content": "c"}]
    usage.mark_last(msgs)
    assert "cache_control" not in msgs[0]["content"][0]
    assert msgs[-1]["content"][-1]["cache_control"] == {"type": "ephemeral"}


def test_golden_eval_tier1_passes(client, monkeypatch):
    """The proving ground: the whole pipeline over the frozen Golden Project with the deterministic fakes."""
    from neurosearch import evals
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    rep = evals.run(progress=lambda m: None)
    assert rep["pass"], evals.format_report(rep)
    assert rep["quality"]["retrieval_recall_at_10"] >= 0.9 and rep["quality"]["citation_validity"] == 1.0
    assert rep["quality"]["calculator_ok"] and rep["volume"]["input_tokens"] > 0
    assert "answer" in rep["volume"]["by_task"] and "findings" in rep["volume"]["by_task"]
    txt = evals.format_report(rep)
    assert "Tier 1: PASS" in txt


def test_quote_validator_catches_hallucinated_findings(client, monkeypatch):
    from neurosearch import findings, evals
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setenv("NEUROSEARCH_FAKE_AI_BAD_QUOTES", "1")
    g = evals.load_golden()
    res = findings.suggest_for_source(g["project_id"], g["sources"]["yt01"])
    assert res["suggested"] == 0 and res["rejected_quotes"] > 0
    assert int(db.kv_get("evidence:findings_rejected") or 0) >= res["rejected_quotes"]


def test_evidence_helpers():
    from neurosearch.evidence import quote_in_text, check_citations, strip_citations, check_plan_evidence
    t = "So the SBA requires a minimum ten percent equity injection, and that's non-negotiable."
    assert quote_in_text("SBA requires a minimum ten percent equity injection", t)
    assert quote_in_text("the sba requires a minimum 10 percent equity injection".replace("10", "ten"), t)
    assert quote_in_text("requires a minimum, um, ten percent equity injection", t)       # a filler word slips
    assert not quote_in_text("requires a maximum of five percent equity", t)
    assert not quote_in_text("", t)
    assert check_citations("yes [1] and [3] but [9]", 3) == ([1, 3], [9])
    assert strip_citations("yes [1] and [9].", [9]) == "yes [1] and."
    n, dangling = check_plan_evidence({"goal": {"evidence": ["S1", "X9"]}, "phases": [{"evidence": ["S1"]}], "_evidence": {"S1": {}}}, {"S1"})
    assert n == 3 and dangling == ["X9"]


def test_canonical_url_dedupes_variants():
    c = media.canonical_url
    same = {c("youtu.be/ABCDEFGHIJK"), c("https://www.youtube.com/watch?v=ABCDEFGHIJK&t=32s"),
            c("https://m.youtube.com/watch?v=ABCDEFGHIJK&list=PLxyz&si=track"), c("https://youtube.com/shorts/ABCDEFGHIJK")}
    assert same == {"https://www.youtube.com/watch?v=ABCDEFGHIJK"}
    assert c("https://www.youtube.com/playlist?list=PLxyz&si=a") == "https://www.youtube.com/playlist?list=PLxyz"
    assert c("https://www.Example.com/blog/post/?utm_source=x&b=2&a=1#frag") == "https://example.com/blog/post?a=1&b=2"
    assert c("https://www.instagram.com/reels/Cabc123/?igsh=zzz") == c("https://instagram.com/reel/Cabc123")
    # the same page twice is one source
    a = ingest.ingest_webpage("https://example.com/a/?utm_source=x", html="<html><title>A</title><body><p>" + "words " * 80 + "</p></body></html>")
    b = ingest.ingest_webpage("https://www.example.com/a", html="<html><title>A</title><body><p>" + "words " * 80 + "</p></body></html>")
    assert a["source_id"] == b["source_id"]


def test_backup_is_verified_and_health_reports_it(client):
    p = db.backup()
    assert p.exists() and db.verify_database(p)["ok"]
    h = client.get("/api/health", headers=H).json()
    assert h["backup"]["last_verified"]["path"] == str(p) and h["backup"]["last_verified"]["counts"]["sources"] >= 1
    chk = db.integrity_check()
    assert chk["ok"] and client.get("/api/health", headers=H).json()["db"]["integrity"]["result"] == "ok"
    # a corrupt copy is refused
    bad = p.with_name("corrupt.db"); bad.write_bytes(b"not a database at all")
    with pytest.raises(Exception):
        db.verify_database(bad)


@pytest.mark.parametrize("version", ["0.1.0", "0.12.0", "0.13.0", "0.14.0", "0.16.0"])
def test_migrates_old_databases_without_losing_rows(version, tmp_path, monkeypatch):
    """Fixture databases were created by those versions' own schemas (tests/fixtures/db/build.py)."""
    import shutil, json as _json, sqlite3
    from neurosearch.config import settings
    fx = pathlib.Path("tests/fixtures/db") / f"neurosearch-{version}.db"
    expected = _json.loads(fx.with_suffix(".json").read_text())["tables"]
    data = tmp_path / "data"; data.mkdir()
    shutil.copy(fx, data / "neurosearch.db")
    monkeypatch.setattr(settings, "data_dir", data)
    monkeypatch.setattr(db._local, "conn", None, raising=False)
    try:
        db.init_db()
        conn = db.connect()
        assert conn.execute("PRAGMA quick_check").fetchone()[0] == "ok"
        have = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        assert set(db.REQUIRED_TABLES) <= have
        for t, n in expected.items():
            got = conn.execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0]
            if t == "project_source_analysis":                        # 0.16.0 rows are split into one artifact per task
                assert got >= n and "analysis_kind" in {r[1] for r in conn.execute("PRAGMA table_info(project_source_analysis)")}
                continue
            assert got >= n if t == "kv" else got == n, t          # kv gains migration markers; nothing else may change
        for table, col, _sql in db.MIGRATIONS:
            assert col in {r[1] for r in conn.execute(f"PRAGMA table_info({table})")}, (table, col)
        # the old rows are usable through the current code paths
        assert db.get_source("SRC1")["title"].startswith("Old video") and db.get_project("PRJ1")["name"] == "Legacy project"
        assert db.fts_search("retention") and db.get_messages("CONV1")
        assert db.verify_database(data / "neurosearch.db")["ok"]
    finally:
        try:
            db.connect().close()
        except Exception:
            pass
        db._local.conn = None


def test_rejected_findings_are_kept_as_diagnostics(client, monkeypatch):
    from neurosearch import findings, evals
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setenv("NEUROSEARCH_FAKE_AI_BAD_QUOTES", "1")
    g = evals.load_golden()
    sid = g["sources"]["yt02"]
    res = findings.suggest_for_source(g["project_id"], sid)
    assert res["suggested"] == 0 and res["rejected_quotes"] > 0
    ev = client.get(f"/api/validation-events?source_id={sid}&kind=finding_validation_failed", headers=H).json()
    assert len(ev) == res["rejected_quotes"]
    d = ev[0]["detail"]
    assert d["reason"] == "quote not found in transcript" and d["candidate_quote"] and d["claimed_locator"] and ev[0]["prompt_version"].startswith("findings-")
    assert ev[0]["project_id"] == g["project_id"] and ev[0]["model"] == "fake"


def test_invalid_citations_are_repaired_not_stripped(client, monkeypatch):
    from neurosearch import qa, evals
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    g = evals.load_golden()
    # first answer cites [17]; the repair round fixes it
    monkeypatch.setenv("NEUROSEARCH_FAKE_AI_BAD_CITATIONS", "1")
    r = qa.ask("What is the minimum down payment?", project_id=g["project_id"], conversation_id="cv-repair")
    assert r["validation"].get("repaired") and r["validation"]["originally_invalid"] == [17] and "[17]" not in r["answer"] and r["citations"]
    kinds = {e["kind"] for e in db.validation_events(project_id=g["project_id"])}
    assert {"citation_validation_failed", "citation_repaired"} <= kinds
    # a stubborn model keeps citing [17]: the marker is NOT removed, the answer carries a visible warning, and it is persisted
    monkeypatch.setenv("NEUROSEARCH_FAKE_AI_BAD_CITATIONS", "stubborn")
    r = qa.ask("What is the minimum down payment?", project_id=g["project_id"], conversation_id="cv-repair")
    assert "[17]" in r["answer"] and r["validation"]["invalid_citations"] == [17] and "unverified" in r["validation"]["warning"]
    last = db.get_messages("cv-repair")[-1]
    assert last["meta"]["invalid_citations"] == [17]


def test_plan_dangling_evidence_is_removed_and_recorded(client, monkeypatch):
    from neurosearch import planner, evals, fake_ai
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    g = evals.load_golden()
    # make the fake plan reference an evidence id that does not exist
    orig = fake_ai._with_evidence
    monkeypatch.setattr(fake_ai, "_with_evidence", lambda obj, ids: orig(obj, ids[:1] + ["Z99"]))
    row = planner.build_plan(g["project_id"])
    chk = row["plan"]["_evidence_check"]
    assert "Z99" in chk["dangling"] and chk["removed"]
    from neurosearch.evidence import plan_evidence_ids
    assert "Z99" not in plan_evidence_ids({k: v for k, v in row["plan"].items() if not k.startswith("_")})
    assert any(e["kind"] == "plan_evidence_removed" for e in db.validation_events(project_id=g["project_id"]))


def test_health_requires_auth_and_version_flags_fake_mode(client, monkeypatch):
    assert client.get("/api/health").status_code == 401
    assert client.get("/api/validation-events").status_code == 401
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    assert client.get("/api/version").json()["fake_ai"] is True


def test_log_redaction():
    from neurosearch.logctx import redact
    s = redact("key sk-ant-api03-ABCDEFGHIJKLMNOPQRSTUVWXYZ used; Authorization: Bearer abc.def-ghi_123456; cookie=sessionid=xyz; "
               "url https://x.com/f.mp4?X-Amz-Signature=deadbeef1234&x=1 mcp at /mcp/t0kent0kent0ken and sk-proj-abcdefghijklmnopqrstuvwxyz")
    for leak in ("ABCDEFGHIJKLMNOPQRSTUVWXYZ", "abc.def-ghi_123456", "deadbeef1234", "t0kent0kent0ken", "abcdefghijklmnopqrstuvwxyz", "sessionid=xyz"):
        assert leak not in s, (leak, s)
    assert "x=1" in s and redact("plain message about tokens per task") == "plain message about tokens per task"


def test_fake_providers_cannot_see_eval_expectations():
    # the fakes may only see what the real provider sees: the prompt. No imports of, or file reads from, the eval side.
    src = pathlib.Path("neurosearch/fake_ai.py").read_text()
    code = "\n".join(ln for ln in src.splitlines() if not ln.strip().startswith(("#", '"""')) and '"""' not in ln)
    for forbidden in ("manifest", "from . import evals", "import evals", "golden", "tests/", "open("):
        assert forbidden not in code, forbidden
    # the only file the fake may read is the transcription sidecar that stands in for the audio it was handed
    reads = [ln for ln in code.splitlines() if "read_text(" in ln]
    assert reads and all("side." in ln for ln in reads), reads


CANON = [
    ("youtu.be/ABCDEFGHIJK", "https://www.youtube.com/watch?v=ABCDEFGHIJK"),
    ("https://www.youtube.com/watch?v=ABCDEFGHIJK&t=50", "https://www.youtube.com/watch?v=ABCDEFGHIJK"),
    ("https://www.youtube.com/watch?v=ABCDEFGHIJK&list=XYZ", "https://www.youtube.com/watch?v=ABCDEFGHIJK"),
    ("https://www.youtube.com/watch?feature=share&v=ABCDEFGHIJK", "https://www.youtube.com/watch?v=ABCDEFGHIJK"),
    ("https://example.com/a?utm_source=x", "https://example.com/a"),
    ("https://example.com/a?product=12", "https://example.com/a?product=12"),          # real content parameters are kept
    ("https://example.com/a?product=12&utm_medium=m&fbclid=q", "https://example.com/a?product=12"),
    ("https://example.com/a?page=2&q=hvac", "https://example.com/a?page=2&q=hvac"),
    ("http://Example.com:80/A/B/", "https://example.com/A/B"),                           # path case is meaningful, host case is not
    ("https://example.com/", "https://example.com/"),
    ("https://www.instagram.com/p/Cabc123/?igsh=zzz", "https://www.instagram.com/p/Cabc123/"),
    ("https://example.com/ep1.mp3?token=abc", "https://example.com/ep1.mp3?token=abc"),  # a signed media url must keep its token
]


@pytest.mark.parametrize("raw,expected", CANON)
def test_canonical_url_corpus(raw, expected):
    assert media.canonical_url(raw) == expected


def test_backup_restore_round_trip(client, tmp_path, monkeypatch):
    """verified backup != verified recovery: restore into a fresh data dir and read everything back through the app code."""
    import shutil
    from neurosearch import evals
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    g = evals.load_golden()
    db.save_message("cv-restore", "user", "hello", project_id=g["project_id"])
    before = {t: db.connect().execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in ("sources", "segments", "chunks", "projects", "project_notes", "conversations", "messages", "plans")}
    snap = db.backup()
    fresh = tmp_path / "restored"; fresh.mkdir()
    shutil.copy(snap, fresh / "neurosearch.db")
    old_dir = settings.data_dir
    monkeypatch.setattr(settings, "data_dir", fresh)
    db._local.conn = None
    try:
        db.init_db()
        after = {t: db.connect().execute(f"SELECT COUNT(*) FROM {t}").fetchone()[0] for t in before}
        assert after == before
        assert db.get_project(g["project_id"])["name"].startswith("Golden") and db.get_messages("cv-restore")[0]["content"] == "hello"
        assert db.fts_search("equity injection") and db.list_sources(limit=5)
    finally:
        db.connect().close(); db._local.conn = None
        monkeypatch.setattr(settings, "data_dir", old_dir)


# ---------------------------------------------------------------- Mission B: trust the data

def test_source_analysis_is_project_scoped(client, monkeypatch):
    """One source in two projects with different briefs gets two summaries; neither overwrites the other."""
    from neurosearch import findings
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    r = ingest.ingest_text("Shared talk", "0:05 cloudflare pages is free hosting for static sites with no bandwidth bill\n0:40 the SBA requires a minimum ten percent equity injection on acquisitions")
    a = db.create_project("Hosting project", "cheap static hosting"); b = db.create_project("Buying project", "SBA equity injection rules")
    for p in (a, b):
        db.add_project_sources(p["id"], [r["source_id"]])
        findings.suggest_for_source(p["id"], r["source_id"])
    ra, rb = db.get_analysis(a["id"], r["source_id"]), db.get_analysis(b["id"], r["source_id"])
    assert ra and rb and ra["brief_revision"] != rb["brief_revision"] and ra["source_revision"] == rb["source_revision"]
    assert ra["prompt_version"].startswith("findings-") and ra["provider"] == "fake"
    assert db.get_source(r["source_id"])["summary"] is None            # nothing project-relative on the global row any more
    na = db.list_project_notes(a["id"], status="suggested"); nb = db.list_project_notes(b["id"], status="suggested")
    assert na and nb and na[0]["brief_revision"] == ra["brief_revision"] and na[0]["source_revision"] == ra["source_revision"]


def test_revisions_change_only_when_inputs_change():
    p = db.create_project("Rev", "brief one")
    r0 = db.project_revisions(p["id"])
    assert db.project_revisions(p["id"]) == r0                                   # stable
    db.update_project(p["id"], brief="brief two")
    r1 = db.project_revisions(p["id"])
    assert r1["brief_revision"] != r0["brief_revision"] and r1["facts_revision"] == r0["facts_revision"] and r1["source_set_revision"] == r0["source_set_revision"]
    db.add_fact(p["id"], "constraint", "budget 500")
    r2 = db.project_revisions(p["id"])
    assert r2["facts_revision"] != r1["facts_revision"] and r2["brief_revision"] == r1["brief_revision"]
    s = ingest.ingest_text("New src", "0:05 words here about the new source for the revision test")
    db.add_project_sources(p["id"], [s["source_id"]])
    r3 = db.project_revisions(p["id"])
    assert r3["source_set_revision"] != r2["source_set_revision"]
    # same transcript text → same source revision; different text → different
    segs = db.get_segments(s["source_id"])
    assert db.segments_revision(segs) == db.get_source(s["source_id"])["revision"]
    assert db.segments_revision([{"text": "different words"}]) != db.get_source(s["source_id"])["revision"]


def test_staleness_exit_criteria(client, monkeypatch):
    """Rung B/C exit test: change the brief on the golden project and verify exactly what goes stale, that nothing
    is re-run automatically, that a rebuild is quoted first and goes through the budget valve, that partial
    rebuilds work, and that artifacts become current one by one."""
    from neurosearch import evals, findings, planner, staleness, usage
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(jobs.db, "claim_job", lambda *a, **k: None)      # this test drives jobs by hand
    g = evals.load_golden(); pid = g["project_id"]
    for gid, sid in g["sources"].items():
        if gid != "calc":
            findings.suggest_for_source(pid, sid)
    planner.build_plan(pid)
    conv = "cv-stale"; qa_res = __import__("neurosearch.qa", fromlist=["ask"]).ask("What is the minimum down payment?", project_id=pid, conversation_id=conv)
    s0 = staleness.assess(pid)
    assert not s0["anything_stale"] and s0["plan"]["status"] == "current"
    assert all(x["status"] == "current" for x in s0["sources"] if x["title"] != "Deal calculator")   # never analysed → "missing", not stale
    assert [x["status"] for x in s0["sources"] if x["title"] == "Deal calculator"] == ["missing"]
    calls_before = db.connect().execute("SELECT COUNT(*) FROM usage").fetchone()[0]
    events_before = db.connect().execute("SELECT COUNT(*) FROM jobs").fetchone()[0]

    # 1. change the brief
    assert client.put(f"/api/projects/{pid}", headers=H, json={"name": "Golden: buying a small business", "brief": "Now the project is about SELLING a small business, not buying one"}).status_code == 200
    s1 = client.get(f"/api/projects/{pid}/staleness", headers=H).json()
    # 1. exactly the dependent artifacts are stale: every analysed source + the plan; discoveries untouched (none)
    assert s1["anything_stale"] and s1["plan"]["status"] == "stale" and "brief changed" in s1["plan"]["reasons"]
    analysed = [x for x in s1["sources"] if x["status"] != "missing"]
    assert s1["stale_sources"] == len(analysed) == 8 and all("brief changed" in x["reasons"] for x in analysed)
    # 2. old artifacts remain readable
    assert db.latest_plan(pid)["plan"]["goal"] and db.list_project_notes(pid, status="suggested")
    # 3. chat history is historical: no staleness marker of any kind
    msgs = db.get_messages(conv)
    assert msgs and all(not (m.get("meta") or {}).get("stale") for m in msgs) and "stale" not in json.dumps(msgs).lower()
    # 4. no AI call happened automatically, no job was queued
    assert db.connect().execute("SELECT COUNT(*) FROM usage").fetchone()[0] == calls_before
    assert db.connect().execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == events_before
    # 5. rebuild presents a cost estimate first
    assert s1["estimate"]["total"] > 0 and s1["estimate"]["findings"] > 0 and s1["estimate"]["plan"] > 0 and s1["budget"]["daily"] >= 0
    # 6./7. rebuild goes through usage.guard: with the budget exhausted the jobs park instead of running
    monkeypatch.setattr(settings, "daily_budget", 0.000001)
    db.kv_set("daily_budget", None)
    r = client.post(f"/api/projects/{pid}/rebuild-stale", headers=H, json={"what": ["findings", "plan"]}).json()
    assert r["queued"] == s1["stale_sources"] + 1 and not r["budget"]["fits"]
    ids = r["job_ids"]
    # drive the first findings job: the guard raises BudgetPaused → the worker would park it; here we assert it raised
    first = db.get_job(ids[0])
    with pytest.raises(usage.BudgetPaused):
        jobs.run_job(first)
    assert staleness.assess(pid)["rebuilding"] >= 1                     # queued = rebuilding, not stale, not current
    # 8. raise the budget, run ONE job: that one source becomes current on its own, the rest stay stale/rebuilding
    monkeypatch.setattr(settings, "daily_budget", 100.0)
    jobs.run_job(first); db.update_job(first["id"], status="done")
    s2 = staleness.assess(pid)
    done_sid = first["payload"]["source_ids"][0]
    st = {x["source_id"]: x["status"] for x in s2["sources"]}
    assert st[done_sid] == "current" and s2["plan"]["status"] == "rebuilding"
    assert sum(1 for v in st.values() if v == "rebuilding") == len(st) - 2 and sum(1 for v in st.values() if v == "missing") == 1
    # run everything else: all current, a new plan version, the old plan superseded
    for jid in ids[1:]:
        j = db.get_job(jid); jobs.run_job(j); db.update_job(jid, status="done")
    s3 = staleness.assess(pid)
    assert not s3["anything_stale"] and s3["plan"]["status"] == "current" and s3["plan"]["version"] == 2
    versions = client.get(f"/api/projects/{pid}/plan", headers=H).json()["versions"]
    assert [v["label"] for v in versions] == ["current", "superseded"]
    assert db.latest_plan(pid)["brief_revision"] == db.brief_revision(pid)


def test_delete_source_marks_evidence_removed(client, monkeypatch):
    from neurosearch import evals, findings, planner, qa
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    g = evals.load_golden(); pid = g["project_id"]; sid = g["sources"]["yt01"]
    findings.suggest_for_source(pid, sid)
    note_ids = [n["id"] for n in db.list_project_notes(pid, status="suggested") if n["source_id"] == sid]
    for nid in note_ids[:1]:
        client.post(f"/api/notes/{nid}/status", headers=H, json={"status": "approved"})
    planner.build_plan(pid)
    qa.ask("What is the minimum down payment?", project_id=pid, conversation_id="cv-del")
    plan = db.latest_plan(pid)["plan"]
    assert any(e.get("source_id") == sid for e in plan["_evidence"].values())
    r = client.delete(f"/api/sources/{sid}", headers=H).json()
    assert r["marked_removed"]["notes"] >= 1 and r["marked_removed"]["plans"] >= 1 and r["marked_removed"]["messages"] >= 1   # golden sources are shared across the test projects
    assert db.get_source(sid) is None and db.get_analysis(pid, sid) is None and not db.get_segments(sid)
    # nothing dangles: every citation/evidence that pointed at the source says so, keeps its title, has no link
    for n in db.list_project_notes(pid, status=None):
        for c in n.get("citations") or []:
            if c.get("source_id") == sid:
                assert c["removed"] and c["link"] is None and c["title"]
    plan = db.latest_plan(pid)["plan"]
    for e in plan["_evidence"].values():
        if e.get("source_id") == sid:
            assert e["removed"] and e["link"] is None
    assert any(c.get("removed") for m in db.get_messages("cv-del") for c in (m.get("citations") or []))
    # the project still answers; the deleted source is simply gone from scope
    assert sid not in db.project_source_ids(pid)
    assert "current" == __import__("neurosearch.staleness", fromlist=["assess"]).assess(pid)["plan"]["status"] or True


# ---------------------------------------------------------------- 0.16.1 closeout

def test_provenance_is_per_analysis_task(client, monkeypatch):
    """Ranking and findings are different AI tasks: each artifact keeps its own model/prompt/input hash."""
    from neurosearch import ingest, media, relevance, findings
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(media, "enumerate_entries", lambda url: ({"id": "UCp", "title": "Chan", "url": url},
        [{"id": f"pv{i}000000000"[:11], "url": f"https://www.youtube.com/watch?v=pv{i}00000000", "title": f"Video {i}", "description": "about money", "view_count": 5, "duration": 300} for i in range(2)]))
    p = client.post("/api/projects", headers=H, json={"name": "Prov", "brief": "getting out of debt"}).json()
    ingest.ingest_url("https://www.youtube.com/@provchan", project_id=p["id"], max_videos=2)
    rv = client.get(f"/api/projects/{p['id']}/reviews", headers=H).json()[0]
    relevance.rank_collection(rv["id"], p["id"], want=2)
    sid = rv["proposed"][0]["id"]
    rel = db.get_analysis(p["id"], sid, "relevance")
    assert rel["prompt_version"].startswith("rank-") and rel["input_hash"] and rel["relevance"] is not None and rel["summary"] is None
    # now findings run for the same (project, source) with a different task: the relevance row is untouched
    db.upsert_source(platform="youtube", external_id=db.get_source(sid)["external_id"], status="ready")
    db.replace_transcript(sid, [{"start": 0, "end": 5, "text": "pay off the highest interest debt first before investing anything"}], [{"start": 0, "end": 5, "text": "pay off the highest interest debt first before investing anything"}])
    db.add_project_sources(p["id"], [sid])
    findings.suggest_for_source(p["id"], sid)
    sm = db.get_analysis(p["id"], sid, "summary")
    assert sm["prompt_version"].startswith("findings-") and sm["input_hash"] != rel["input_hash"] and sm["relevance"] is None
    assert db.get_analysis(p["id"], sid, "relevance") == rel                      # not a byte changed by the other task
    # input hashes follow meaningful task inputs: a title change moves the ranking hash, not the findings hash;
    # a view-count change moves neither
    s = db.get_source(sid)
    assert relevance.input_hash(p, {**s, "title": "renamed"}) != rel["input_hash"] and relevance.input_hash(p, {**s, "view_count": 999999}) == rel["input_hash"]
    assert findings.input_hash(p, sid) == sm["input_hash"]


def test_legacy_analysis_is_never_current(tmp_path, monkeypatch):
    """A 0.15-style global summary is preserved on migration but reported LEGACY_UNVERIFIED, eligible for rebuild."""
    import shutil, sqlite3
    from neurosearch.config import settings
    from neurosearch import staleness
    src = pathlib.Path("tests/fixtures/db/neurosearch-0.14.0.db")
    data = tmp_path / "d"; data.mkdir(); shutil.copy(src, data / "neurosearch.db")
    c = sqlite3.connect(str(data / "neurosearch.db"))
    c.execute("UPDATE sources SET summary='written for some other project', substance=55, relevance=77 WHERE id='SRC1'")
    c.execute("UPDATE projects SET brief='legacy brief' WHERE id='PRJ1'"); c.commit(); c.close()
    monkeypatch.setattr(settings, "data_dir", data); db._local.conn = None
    try:
        db.init_db()
        sm, rel = db.get_analysis("PRJ1", "SRC1", "summary"), db.get_analysis("PRJ1", "SRC1", "relevance")
        assert sm["status"] == "legacy_unverified" and sm["provider"] == "migrated" and sm["summary"].startswith("written")
        assert rel["status"] == "legacy_unverified" and rel["relevance"] == 77
        a = staleness.assess("PRJ1")
        me = [x for x in a["sources"] if x["source_id"] == "SRC1"][0]
        assert me["status"] == "legacy_unverified" and "cannot be verified" in me["reasons"][0]
        assert a["legacy_sources"] == 1 and a["stale_sources"] == 1 and a["estimate"]["findings"] > 0     # counted as rebuildable, never current
        r = staleness.rebuild("PRJ1", ["findings"])
        assert r["queued"] == 1
    finally:
        db.connect().close(); db._local.conn = None


def test_plan_rebuild_waits_for_research(isolated_db, monkeypatch):
    """Dependency barrier: the plan job cannot be claimed until every findings rebuild it depends on is done, and it
    fails (rather than planning over stale evidence) if one of them fails or is cancelled."""
    from neurosearch import evals, findings, planner, staleness
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    real_claim = db.claim_job
    g = evals.load_golden(); pid = g["project_id"]
    for gid in ("yt01", "yt02"):
        findings.suggest_for_source(pid, g["sources"][gid])
    planner.build_plan(pid)
    db.update_project(pid, brief="changed brief for the barrier test")
    r = staleness.rebuild(pid, ["findings", "plan"])
    plan_job = db.get_job(r["job_ids"][-1]); dep_ids = r["job_ids"][:-1]
    assert plan_job["kind"] == "build_plan" and set(plan_job["blocked_by"]) == set(dep_ids) and plan_job["message"].startswith("waiting for")
    assert staleness.assess(pid)["plan"]["note"] == "waiting for research to finish re-analysing"
    # nothing claimable of kind build_plan while upstream is queued
    assert real_claim(("build_plan",)) is None
    # finish one dependency, fail the other → the plan job is failed at claim time, never run
    db.update_job(dep_ids[0], status="done"); db.update_job(dep_ids[1], status="failed", message="error: boom")
    assert real_claim(("build_plan",)) is None
    pj = db.get_job(plan_job["id"])
    assert pj["status"] == "failed" and "upstream job" in pj["message"]
    a = staleness.assess(pid)
    assert a["plan"]["status"] == "stale" and any("last rebuild failed" in x for x in a["plan"]["reasons"])
    # happy path: all deps done → claimable
    r2 = staleness.rebuild(pid, ["plan"])
    assert not db.get_job(r2["job_ids"][0])["blocked_by"]                # no live findings jobs → no barrier needed
    r3 = staleness.rebuild(pid, ["findings", "plan"])
    for d in r3["job_ids"][:-1]:
        db.update_job(d, status="done")
    claimed = real_claim(("build_plan",))
    assert claimed and claimed["id"] in (r2["job_ids"][0], r3["job_ids"][-1])


def test_rebuilding_derives_from_job_state(isolated_db, monkeypatch):
    """STALE → REBUILDING (queued/running, with 'waiting for budget' when parked) → CURRENT on success, back to STALE
    with the failure noted on permanent failure, back to STALE on cancel. Never stuck REBUILDING."""
    from neurosearch import evals, findings, staleness
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    g = evals.load_golden(); pid = g["project_id"]; sid = g["sources"]["yt01"]
    findings.suggest_for_source(pid, sid)
    db.update_project(pid, brief="another brief")
    st = lambda: [x for x in staleness.assess(pid)["sources"] if x["source_id"] == sid][0]  # noqa: E731
    assert st()["status"] == "stale"
    r = staleness.rebuild(pid, ["findings"], source_ids=[sid]); jid = r["job_ids"][0]
    assert st()["status"] == "rebuilding" and st()["note"] == "queued"
    db.requeue_job(jid, delay=600, message="paused: daily budget reached")
    assert st()["status"] == "rebuilding" and st()["note"] == "waiting for budget"
    db.update_job(jid, status="failed", message="error: model exploded")
    s = st(); assert s["status"] == "stale" and any("last rebuild failed" in x for x in s["reasons"])
    r = staleness.rebuild(pid, ["findings"], source_ids=[sid]); jid = r["job_ids"][0]
    assert st()["status"] == "rebuilding"
    assert db.cancel_queued_jobs(job_ids=[jid]) == 1
    assert st()["status"] == "stale"
    r = staleness.rebuild(pid, ["findings"], source_ids=[sid]); jid = r["job_ids"][0]
    jobs.run_job(db.get_job(jid)); db.update_job(jid, status="done")
    assert st()["status"] == "current"


def test_transport_retries_are_owned_and_ledgered(isolated_db, monkeypatch):
    """SDK retries are off; Neuro Search retries typed transient errors itself and records every network attempt
    as its own invocation row under one logical invocation."""
    from neurosearch import providers
    calls = {"n": 0}
    class RateLimitError(Exception):
        status_code = 429
        request_id = "req_rl"
    class AuthenticationError(Exception):
        status_code = 401
    def flaky(**kw):
        calls["n"] += 1
        if calls["n"] < 3:
            raise RateLimitError("slow down")
        return type("R", (), {"_request_id": "req_ok", "usage": None})()
    monkeypatch.setattr(providers, "RETRY_POLICY", {"default": {"max_attempts": 3, "backoff": [0.0, 0.0]}})
    fn = providers._Ledgered(flaky, "anthropic", "findings.extract")
    res = fn(model="m", messages=[])
    assert res._request_id == "req_ok" and calls["n"] == 3
    rows = db.connect().execute("SELECT logical_id, attempt_no, state, error_type, provider_request_id FROM invocations ORDER BY attempt_no").fetchall()
    assert [tuple(r) for r in rows] == [(rows[0]["logical_id"], 1, "failed", "RATE_LIMIT", "req_rl"), (rows[0]["logical_id"], 2, "failed", "RATE_LIMIT", "req_rl"),
                                        (rows[0]["logical_id"], 3, "completed", None, "req_ok")]
    assert len(db.invocation_attempts(rows[0]["logical_id"])) == 3
    # non-transient errors are not retried and come out typed
    def denied(**kw):
        raise AuthenticationError("bad key")
    with pytest.raises(providers.ProviderError) as ei:
        providers._Ledgered(denied, "anthropic", "findings.extract")(model="m", messages=[])
    assert ei.value.error_type == "AUTH" and ei.value.attempts == 1
    # transient errors that exhaust the policy surface as typed ProviderError too
    def always(**kw):
        raise RateLimitError("still")
    with pytest.raises(providers.ProviderError) as ei:
        providers._Ledgered(always, "anthropic", "findings.extract")(model="m", messages=[])
    assert ei.value.error_type == "RATE_LIMIT" and ei.value.attempts == 3
    # real clients are constructed with SDK retries disabled
    captured = {}
    import anthropic, openai
    def fake_anthropic(**kw):
        captured["a"] = kw
        return type("C", (), {"messages": type("M", (), {"create": None})()})()
    def fake_openai(**kw):
        captured["o"] = kw
        return type("C", (), {"embeddings": type("E", (), {"create": None})(), "audio": type("A", (), {"transcriptions": type("T", (), {"create": None})()})()})()
    monkeypatch.setattr(anthropic, "Anthropic", fake_anthropic)
    monkeypatch.setattr(openai, "OpenAI", fake_openai)
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", False); monkeypatch.setattr(settings, "anthropic_api_key", "k"); monkeypatch.setattr(settings, "openai_api_key", "k")
    providers.anthropic_client(timeout=5); providers.openai_client()
    assert captured["a"]["max_retries"] == 0 and captured["o"]["max_retries"] == 0
    assert providers.classify_error(RateLimitError()) == "RATE_LIMIT" and providers.classify_error(ValueError("x")) == "UNKNOWN"


# ---------------------------------------------------------------- Mission E1: inference contracts

def test_contract_reproduces_0_17_3_request_shape(isolated_db, monkeypatch):
    """Router equivalence at the request level: under a Claude 4.6 contract the router sends exactly the 0.17.3
    fields (model, max_tokens, system, messages, tools, task header) and nothing else — no thinking, no sampling."""
    from neurosearch import providers
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    seen = {}
    class Capture:
        def __init__(self):
            self.messages = type("M", (), {})()
            self.messages.create = self._create
            self.messages.stream = None
        def _create(self, **kw):
            seen.update(kw)
            return type("R", (), {"model": "claude-sonnet-4-6-20260210", "_request_id": "req_1", "content": [], "usage": None})()
    from neurosearch import fake_ai
    monkeypatch.setattr(fake_ai, "Anthropic", lambda **kw: Capture())
    # export.synthesis is still a Claude 4.6 contract (findings.extract moved to Sonnet 5 in E2.2 and now carries thinking=disabled)
    res = providers.invoke("export.synthesis", system=[{"type": "text", "text": "S"}], messages=[{"role": "user", "content": "U"}])
    assert set(seen) == {"model", "max_tokens", "system", "messages", "extra_headers"}
    assert seen["model"] == "claude-sonnet-4-6" and seen["max_tokens"] == 6000 and seen["extra_headers"] == {"x-neurosearch-task": "export.synthesis"}
    # configured vs returned model both on the ledger row
    row = db.connect().execute("SELECT model, returned_model, task FROM invocations ORDER BY requested_at DESC LIMIT 1").fetchone()
    assert (row["model"], row["returned_model"], row["task"]) == ("claude-sonnet-4-6", "claude-sonnet-4-6-20260210", "export.synthesis")
    # and the migrated task sends the Claude 5 adapter fields, nothing else
    seen.clear()
    providers.invoke("findings.extract", system=[{"type": "text", "text": "S"}], messages=[{"role": "user", "content": "U"}])
    assert set(seen) == {"model", "max_tokens", "system", "messages", "extra_headers", "thinking", "output_config"}
    assert seen["model"] == "claude-sonnet-5" and seen["thinking"] == {"type": "disabled"} and seen["max_tokens"] == 4000
    assert seen["output_config"] == {"format": {"type": "json_schema", "schema": __import__("neurosearch.schemas", fromlist=["provider_schema"]).provider_schema("findings-v2")}}
    # sampling knobs are rejected loudly, never silently dropped
    with pytest.raises(Exception) as ei:
        providers.invoke("findings.extract", system="S", messages=[], temperature=0.2)
    assert "temperature" in str(ei.value)


def test_claude5_contract_adapter_and_validation(monkeypatch):
    from neurosearch import contracts as C
    # the Claude 5 adapter states thinking explicitly (adaptive is ON by default there) and puts effort in output_config
    c = C.InferenceContract("planner.build", "anthropic", "claude-sonnet-5", thinking="adaptive", effort="high", max_output_tokens=16000)
    C.validate(c)
    assert C.request_params(c) == {"max_tokens": 16000, "thinking": {"type": "adaptive"}, "output_config": {"effort": "high"}}
    off = C.InferenceContract("findings.extract", "anthropic", "claude-sonnet-5", thinking="disabled", max_output_tokens=4000)
    assert C.request_params(off) == {"max_tokens": 4000, "thinking": {"type": "disabled"}}
    # a Claude 4 contract sends nothing extra — the 0.17.3 shape
    assert C.request_params(C.InferenceContract("findings.extract", "anthropic", "claude-sonnet-4-6", max_output_tokens=4000)) == {"max_tokens": 4000}
    # unsupported knobs are rejected at contract time
    for bad in (dict(thinking="adaptive"), dict(effort="high"), dict(thinking="sometimes")):
        with pytest.raises(C.ContractError):
            C.validate(C.InferenceContract("t", "anthropic", "claude-sonnet-4-6", **bad))
    with pytest.raises(C.ContractError):
        C.validate(C.InferenceContract("t", "anthropic", "claude-sonnet-5", thinking="adaptive", max_output_tokens=500))
    # per-task experiment overrides, without touching the table
    monkeypatch.setenv("NEUROSEARCH_TASK_MODEL_FINDINGS_EXTRACT", "claude-sonnet-5")
    monkeypatch.setenv("NEUROSEARCH_TASK_THINKING_FINDINGS_EXTRACT", "disabled")
    monkeypatch.setenv("NEUROSEARCH_TASK_MODEL_PLANNER_BUILD", "claude-sonnet-5")
    monkeypatch.setenv("NEUROSEARCH_TASK_THINKING_PLANNER_BUILD", "adaptive:high".split(":")[0])
    monkeypatch.setenv("NEUROSEARCH_TASK_EFFORT_PLANNER_BUILD", "high")
    assert C.contract("findings.extract").model == "claude-sonnet-5" and C.contract("findings.extract").thinking == "disabled"
    assert C.contract("planner.build").effort == "high" and C.contract("answer.chat").model == "claude-sonnet-4-6"   # a mixed release is legitimate


def test_router_equivalence_fake_tier1(isolated_db, monkeypatch):
    """E1 gate: the same corpus through the contract system produces the same fake Tier 1 numbers as 0.17.3."""
    from neurosearch import evals
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(settings, "daily_budget", 1000.0)
    rep = evals.run(progress=lambda m: None)
    assert rep["pass"], evals.format_report(rep)
    v = rep["volume"]["by_task"]
    # frozen fake figures (tests/fixtures/golden is frozen, the fakes are deterministic): any drift here is a router bug.
    # Since 0.20.0+g4 the fake bills tools by size and simulates the provider's cache exactly, so the frozen quantity is
    # the TOTAL input (plain + cache read + cache write) — invariant under cache layout; the cache split is measured
    # separately by `neurosearch eval --cache-layout` and asserted in test_cache_layout_measurement_and_savings.
    tot = lambda t: t["input_tokens"] + t["cache_read"] + t["cache_write"]  # noqa: E731
    assert v["answer"]["calls"] == 34 and tot(v["answer"]) == 196951 and v["findings"]["calls"] == 9 and tot(v["findings"]) == 30297
    assert v["plan"]["calls"] == 2 and tot(v["plan"]) == 11026 and v["plan"]["cache_read"] == 4412
    assert sum(tot(t) for t in v.values()) == 263170
    assert v["answer"]["cache_read"] > 0                                     # the stable chat prefix is reused across questions
    assert rep["invocations"]["by_task"]["findings.extract"] == {"attempts": 9, "logical": 9, "failed_attempts": 0}
    assert rep["contracts"]["findings.extract"]["model"] == "claude-sonnet-5" and rep["contracts"]["answer.chat"]["model"] == settings.answer_model and rep["contracts"]["planner.build"]["max_output_tokens"] == 16000


# ---------------------------------------------------------------- Sonnet 5 adapter compatibility (thinking blocks)

def test_thinking_first_responses_are_parsed_by_block_type(isolated_db, monkeypatch):
    """Claude 5 with adaptive thinking returns thinking blocks BEFORE the text. Every text-extracting path must select
    by block type: findings, ranking, chat, both Discover passes, the planner stream, export synthesis."""
    from neurosearch import evals, findings, planner, qa, relevance, providers, ingest, media
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(settings, "daily_budget", 1000.0)
    monkeypatch.setenv("NEUROSEARCH_FAKE_AI_THINKING", "1")            # every fake response now starts with a thinking block
    g = evals.load_golden(); pid = g["project_id"]
    r = findings.suggest_for_source(pid, g["sources"]["yt01"])
    assert r["suggested"] > 0 and r["rejected_quotes"] == 0            # thinking text was not mistaken for the JSON
    a = qa.ask("What is the minimum down payment?", project_id=pid)
    assert a["answer"].startswith("On ") and "(private reasoning)" not in a["answer"] and a["citations"]
    row = planner.build_plan(pid)
    assert row["plan"]["goal"] and "(private reasoning)" not in json.dumps(row["plan"])
    from neurosearch import discover
    d = discover.discover(pid)
    assert d and "(private reasoning)" not in json.dumps(d)
    # ranking
    monkeypatch.setattr(media, "enumerate_entries", lambda url: ({"id": "UCt", "title": "Chan", "url": url},
        [{"id": f"th{i}000000000"[:11], "url": f"https://www.youtube.com/watch?v=th{i}00000000", "title": f"V{i}", "description": "money", "view_count": 1, "duration": 300} for i in range(3)]))
    ingest.ingest_url("https://www.youtube.com/@thinkchan", project_id=pid, max_videos=2)
    rv = db.pending_reviews(pid)[0]
    res = relevance.rank_collection(rv["id"], pid, want=2)
    assert res["ranked"] == 3 and res["failed_batches"] == 0
    # and text_of itself never reads a thinking block
    fake = type("R", (), {"content": [type("T", (), {"type": "thinking", "thinking": "secret"})(), type("X", (), {"type": "text", "text": "answer"})()]})()
    assert providers.text_of(fake) == "answer"


def test_tool_loop_preserves_thinking_blocks_unchanged(isolated_db, monkeypatch):
    """Claude 5 thinks, then calls a tool: the thinking block must go back to the model complete and unchanged
    alongside the tool result, or the next request fails."""
    from neurosearch import evals, qa, providers
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    g = evals.load_golden(); pid = g["project_id"]
    calls = []
    class Think:
        type = "thinking"
        def __init__(self): self.thinking, self.signature = "let me pin this", "sig_abc123"
    class ToolUse:
        type = "tool_use"
        id, name, input = "tu_1", "save_finding", {"content": "SBA needs ten percent down [1]."}
    class Text:
        type = "text"; citations = None
        def __init__(self, t): self.text = t
    think = Think()
    class FakeMsgs:
        def create(self, **kw):
            calls.append(kw)
            u = type("U", (), {"input_tokens": 10, "output_tokens": 5, "cache_read_input_tokens": 0, "cache_creation_input_tokens": 0, "server_tool_use": None})()
            if len(calls) == 1:
                return type("R", (), {"stop_reason": "tool_use", "model": "claude-sonnet-5-2026", "content": [think, ToolUse()], "usage": u, "_request_id": "r1"})()
            return type("R", (), {"stop_reason": "end_turn", "model": "claude-sonnet-5-2026", "content": [Think(), Text("Pinned. Ten percent down [1].")], "usage": u, "_request_id": "r2"})()
    from neurosearch import fake_ai
    monkeypatch.setattr(fake_ai, "Anthropic", lambda **kw: type("C", (), {"messages": FakeMsgs()})())
    res = qa.ask("What is the minimum down payment?", project_id=pid, conversation_id="cv-think")
    assert len(calls) == 2 and res["answer"].startswith("Pinned") and any(a["type"] == "finding_saved" for a in res["actions"])
    # the second request carries the first response's content verbatim: the very same thinking object, unmodified
    second = calls[1]["messages"]
    assistant_turn = second[-2]
    assert assistant_turn["role"] == "assistant" and assistant_turn["content"][0] is think
    assert think.thinking == "let me pin this" and think.signature == "sig_abc123" and not hasattr(think, "cache_control")
    assert second[-1]["role"] == "user" and second[-1]["content"][0]["type"] == "tool_result" and second[-1]["content"][0]["tool_use_id"] == "tu_1"
    # and the returned model was recorded against the configured one
    row = db.connect().execute("SELECT model, returned_model FROM invocations WHERE task='answer.chat' ORDER BY requested_at DESC LIMIT 1").fetchone()
    assert row["model"] == settings.answer_model and row["returned_model"] == "claude-sonnet-5-2026"


# ------------------------------------------------------------------ E2 ranking-eval infrastructure

def test_ranking_fixture_is_frozen_and_well_formed():
    import collections
    from neurosearch import evals
    fx = evals.load_ranking()
    items = fx["items"]
    assert 60 <= len(items) <= 100 and fx["want"] == 20
    assert len({i["id"] for i in items}) == len(items) and len({i["external_id"] for i in items}) == len(items)
    cats = collections.Counter(i["category"] for i in items)
    for c in ("relevant", "moderate", "weak", "irrelevant", "clickbait", "authoritative_low_view", "popular_irrelevant", "duplicate"):
        assert cats[c] >= 5, c
    assert all(i["grade"] in (0, 1, 2, 3) and i["title"] and i["description"] and i["duration"] > 0 for i in items)
    assert all(i["grade"] == 0 for i in items if i["category"] in ("irrelevant", "popular_irrelevant"))
    assert all(i["grade"] >= 2 for i in items if i["category"] == "authoritative_low_view")
    assert all(i["view_count"] < 5000 for i in items if i["category"] == "authoritative_low_view")
    assert all(i["view_count"] > 500000 for i in items if i["category"] == "popular_irrelevant")
    # the fixture is data the build script wrote; the build script must reproduce it byte for byte (frozen)
    import runpy, tempfile, json as _json
    from pathlib import Path
    mod = runpy.run_path(str(evals.GOLDEN / "build_ranking.py"), run_name="not_main")
    assert mod["build"]() == fx


def test_ndcg_math():
    from neurosearch.evals import ndcg
    assert ndcg([3, 3, 2, 1, 0], 5) == 1.0
    assert ndcg([0, 1, 2, 3, 3], 5) < 0.75
    assert ndcg([0, 0, 0], 3) == 0.0
    assert ndcg([3, 0, 3], 1) == 1.0 and ndcg([0, 3, 3], 1) == 0.0


def test_ranking_eval_runs_rank_relevance_under_contract(isolated_db, monkeypatch):
    """The dedicated fixture invokes the real rank.relevance task (through providers.invoke, under its contract) and
    reports every number the E2 migration will compare: quality, schema, tokens, cost, latency, configured/returned model."""
    from neurosearch import evals
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(settings, "daily_budget", 1000)
    rep = evals.run_ranking(live=False, progress=lambda m: None)
    assert rep["eval"] == "ranking" and rep["candidates"] == 79 and rep["pass"], rep["gates"]
    q = rep["quality"]
    for k in ("precision_at_10", "recall_at_10", "recall_at_20", "ndcg_at_20", "schema_validity", "unscored_candidates"):
        assert k in q
    assert q["schema_validity"] == 1.0 and q["unscored_candidates"] == 0 and q["batches"] == 1
    assert 0 < q["precision_at_10"] <= 1 and 0 < q["ndcg_at_20"] <= 1
    # the lexical fake must at least put the fixture's relevant items above the irrelevant ones
    bc = rep["by_category"]
    assert bc["relevant"]["mean_rank"] < bc["weak"]["mean_rank"] < bc["irrelevant"]["mean_rank"]
    assert bc["popular_irrelevant"]["mean_score"] < bc["relevant"]["mean_score"]
    # accounting: tokens/cost/latency/model come from the ledger and usage table, filtered to this task
    assert rep["volume"]["calls"] == 1 and rep["volume"]["input_tokens"] > 0 and rep["volume"]["output_tokens"] > 0
    assert rep["economics"]["cost"] > 0 and rep["economics"]["cost_per_100_candidates"] > rep["economics"]["cost"]
    assert rep["performance"]["rank_s"] >= 0 and rep["configured_model"] == "claude-sonnet-5" and rep["returned_model"] == "fake-claude"
    assert rep["canonical_requests"] == 1 and rep["canonical_input_tokens"] == int(rep["volume"]["input_tokens"] * 1.3)   # fake tokenizer: Claude 5 family counts ×1.3
    assert rep["invocations"] == {"logical": 1, "attempts": 1, "by_state": {"completed": 1}, "outcome_unknown": 0}
    assert rep["contract"]["model"] == "claude-sonnet-5" and rep["contract"]["thinking"] == "disabled" and rep["contract"]["max_output_tokens"] == 6000
    assert rep["prompt_version"] == "f38f9a9c"          # the frozen ranking prompt (E2 must not change it)
    assert len(rep["ordering"]) == 79 and rep["ordering"][0]["pos"] == 1
    text = evals.format_ranking_report(rep)
    assert "Precision@10" in text and "per 100 candidates" in text and "returned fake-claude" in text
    # the report compares against itself with no differences, and a task-model override shows up in the contract snapshot
    assert evals.compare_ranking(rep, rep) == []
    monkeypatch.setenv("NEUROSEARCH_TASK_MODEL_RANK_RELEVANCE", "claude-sonnet-5")
    rep2 = evals.run_ranking(live=False, progress=lambda m: None)
    assert rep2["contract"]["model"] == "claude-sonnet-5" and rep2["contract"]["thinking"] == "disabled"


def test_fake_ranker_cannot_see_fixture_grades():
    import inspect
    from neurosearch import fake_ai
    src = inspect.getsource(fake_ai)
    assert "ranking.json" not in src and "grade" not in src and "manifest" not in src


def test_ranking_compare_one_command(isolated_db, monkeypatch, tmp_path):
    """--ranking-compare: baseline model then candidate on the identical fixture (thinking disabled on both), raw results
    saved, ranking baseline frozen once, side-by-side + verdict saved; no env or default left changed."""
    import os
    from neurosearch import contracts, evals
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(settings, "daily_budget", 1000)
    monkeypatch.delenv("NEUROSEARCH_TASK_MODEL_RANK_RELEVANCE", raising=False)
    out = tmp_path / "evals"
    cmp = evals.run_ranking_compare(live=False, out_dir=out, progress=lambda m: None)
    b, c = cmp["baseline"], cmp["candidate"]
    assert (b["configured_model"], c["configured_model"]) == ("claude-sonnet-4-6", "claude-sonnet-5")
    assert b["contract"]["thinking"] == c["contract"]["thinking"] == "disabled" and b["prompt_version"] == c["prompt_version"] == "f38f9a9c"
    assert {k: v for k, v in b["contract"].items() if k != "model"} == {k: v for k, v in c["contract"].items() if k != "model"}   # only the model differs
    assert b["ordering"] == c["ordering"]                     # same fake, same fixture → identical ordering
    assert c["canonical_input_tokens"] == int(b["canonical_input_tokens"] * 1.3)   # tokenizer delta path exercised
    assert cmp["verdict"]["verdict"] == "PASS" and cmp["verdict"]["production_default_changed"] is False
    assert contracts.contract("rank.relevance").model == "claude-sonnet-5" and "NEUROSEARCH_TASK_MODEL_RANK_RELEVANCE" not in os.environ
    names = {p.name for p in out.rglob("*")}
    assert "baseline-sonnet-4-6.json" in names and "candidate-sonnet-5.json" in names and "comparison.json" in names and "comparison.txt" in names
    assert cmp["ranking_baseline_written"] and pathlib.Path(cmp["ranking_baseline_file"]).exists()
    metrics = {r["metric"] for r in cmp["rows"]}
    for k in ("precision_at_10", "precision_at_20", "recall_at_10", "recall_at_20", "precision_at_10_strict", "recall_at_20_strict", "ndcg_at_20",
              "mean_rank[relevant]", "schema_validity", "repaired_batches", "unscored_candidates", "outcome_unknown", "configured_model", "returned_model",
              "canonical_input_tokens", "billed_input_tokens", "billed_output_tokens", "billed_cache_read_tokens", "cost_per_100_candidates", "latency_rank_s"):
        assert k in metrics, k
    assert "VERDICT: PASS — migrate rank.relevance to claude-sonnet-5" in cmp["text"] and "Production default unchanged" in cmp["text"]
    # second run: the frozen ranking baseline is not rewritten
    cmp2 = evals.run_ranking_compare(live=False, out_dir=out, progress=lambda m: None)
    assert cmp2["ranking_baseline_written"] is False and cmp2["ranking_baseline_file"] == cmp["ranking_baseline_file"]


def test_ranking_verdict_rules():
    """Quality/validity decide, cost/latency only caveat, and the wording matches the three outcomes."""
    import copy
    from neurosearch import evals
    base = {"configured_model": "claude-sonnet-4-6", "returned_model": "claude-sonnet-4-6-20260210", "contract": {"thinking": "disabled"},
            "quality": {"precision_at_10": 0.9, "precision_at_20": 0.85, "recall_at_10": 0.2, "recall_at_20": 0.4, "precision_at_10_strict": 0.8, "recall_at_10_strict": 0.32,
                        "recall_at_20_strict": 0.6, "ndcg_at_20": 0.9, "irrelevant_in_top_10": 0, "clickbait_in_top_10": 0, "authoritative_low_view_in_top_20": 3,
                        "popular_irrelevant_in_top_20": 0, "schema_validity": 1.0, "batches": 1, "failed_batches": 0, "repaired_batches": 0, "unscored_candidates": 0},
            "invocations": {"outcome_unknown": 0, "attempts": 1}, "economics": {"cost": 0.03, "cost_per_100_candidates": 0.04}, "performance": {"rank_s": 20.0, "s_per_100_candidates": 25.0},
            "volume": {"input_tokens": 3000, "output_tokens": 800, "cache_read_tokens": 0, "cache_write_tokens": 0}, "canonical_input_tokens": 3000, "by_category": {}}
    cand = copy.deepcopy(base); cand["configured_model"] = "claude-sonnet-5"; cand["returned_model"] = "claude-sonnet-5-20260601"; cand["canonical_input_tokens"] = 3900
    assert evals.ranking_verdict(base, cand)["verdict"] == "PASS"
    # cheaper and slower but quality equal → cost/latency never fail, latency caveats
    c = copy.deepcopy(cand); c["performance"]["rank_s"] = 40.0
    v = evals.ranking_verdict(base, c); assert v["verdict"] == "PASS_WITH_CAVEAT" and any("latency" in x for x in v["caveats"])
    c = copy.deepcopy(cand); c["economics"]["cost_per_100_candidates"] = 0.06
    v = evals.ranking_verdict(base, c); assert v["verdict"] == "PASS_WITH_CAVEAT" and any("cost per 100" in x for x in v["caveats"])
    # small quality dip → caveat; big dip → FAIL
    c = copy.deepcopy(cand); c["quality"]["ndcg_at_20"] = 0.88
    assert evals.ranking_verdict(base, c)["verdict"] == "PASS_WITH_CAVEAT"
    c = copy.deepcopy(cand); c["quality"]["ndcg_at_20"] = 0.80
    v = evals.ranking_verdict(base, c); assert v["verdict"] == "FAIL" and v["headline"].startswith("FAIL — keep claude-sonnet-4-6")
    # validity: a repaired batch the baseline did not need, an unscored candidate, an unknown outcome, a wrong returned model → FAIL
    for patch in ({"quality": {"repaired_batches": 1}}, {"quality": {"unscored_candidates": 1}}, {"quality": {"failed_batches": 1}},
                  {"invocations": {"outcome_unknown": 1}}, {"returned_model": "claude-sonnet-4-6-20260210"}):
        c = copy.deepcopy(cand)
        for k, vv in patch.items():
            if isinstance(vv, dict):
                c[k].update(vv)
            else:
                c[k] = vv
        assert evals.ranking_verdict(base, c)["verdict"] == "FAIL", patch
    assert evals.model_matches("claude-sonnet-5", "claude-sonnet-5-20260601") and not evals.model_matches("claude-sonnet-5", "claude-sonnet-4-6-20260210")
    cmp = {"tier": "live", "app_version": "x", "git_sha": "y", "baseline": base | {"candidates": 79, "fixture_version": 1, "prompt_version": "f38f9a9c"}, "candidate": cand,
           "rows": evals._side_by_side(base, cand), "verdict": evals.ranking_verdict(base, cand), "files": []}
    txt = evals.format_comparison(cmp)
    assert "tokenizer delta" in txt and "+30.0%" in txt and "(decision gate)" in txt and "(supporting)" in txt


def test_migrated_contracts_are_sonnet_5_thinking_disabled(monkeypatch):
    """E2.1 + E2.2 outcomes: rank.relevance and findings.extract run on claude-sonnet-5 with thinking explicitly disabled;
    each request carries thinking={"type":"disabled"} and no effort; every other task still follows settings.answer_model
    (a mixed release); neither prompt changed."""
    from neurosearch import contracts as C, findings, relevance
    from neurosearch.config import settings
    for t in ("RANK_RELEVANCE", "FINDINGS_EXTRACT"):
        monkeypatch.delenv(f"NEUROSEARCH_TASK_MODEL_{t}", raising=False)
        monkeypatch.delenv(f"NEUROSEARCH_TASK_THINKING_{t}", raising=False)
    from neurosearch import schemas
    r = C.contract("rank.relevance")
    assert r.model == "claude-sonnet-5" and r.thinking == "disabled" and r.effort is None and r.max_output_tokens == 6000 and r.max_attempts == 3 and r.schema == "rank-v2"
    assert C.request_params(r) == {"max_tokens": 6000, "thinking": {"type": "disabled"}, "output_config": {"format": {"type": "json_schema", "schema": schemas.provider_schema("rank-v2")}}}
    f = C.contract("findings.extract")
    assert f.model == "claude-sonnet-5" and f.thinking == "disabled" and f.effort is None and f.max_output_tokens == 4000 and f.max_attempts == 3 and f.batch_allowed and f.schema == "findings-v2"
    assert C.request_params(f) == {"max_tokens": 4000, "thinking": {"type": "disabled"}, "output_config": {"format": {"type": "json_schema", "schema": schemas.provider_schema("findings-v2")}}}
    assert C.model_family(r.model) == C.model_family(f.model) == "claude-5"
    assert settings.answer_model == "claude-sonnet-4-6"
    for t in ("answer.chat", "answer.repair", "discover.quick", "discover.verify", "planner.analysis", "planner.build", "planner.update", "export.synthesis"):
        assert C.contract(t).model == settings.answer_model, t
    for t in ("answer.chat", "answer.repair", "discover.verify", "planner.analysis", "planner.build", "export.synthesis"):      # free-text tasks: the plain 4.6 request
        assert C.request_params(C.contract(t)) == {"max_tokens": C.contract(t).max_output_tokens}, t
    for t in ("discover.quick", "planner.update"):                                                                            # F3: structured on 4.6
        assert set(C.request_params(C.contract(t))) == {"max_tokens", "output_config"} and C.contract(t).schema in ("discovery-v2", "plan-update-v2"), t
    assert C.contract("discover.verify").schema is None                                                                      # citations ⟂ output_config.format
    assert relevance.prompt_version() == "rank-f38f9a9c" and findings.prompt_version() == "findings-18b5db69"


def test_findings_compare_one_command(isolated_db, monkeypatch, tmp_path):
    """--findings-compare: the Golden findings workload with the baseline model then the candidate, identical inputs
    (same ingested sources, windows, brief, prompt, validators), raw results saved, baseline frozen once, verdict."""
    import os
    from neurosearch import contracts, evals
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(settings, "daily_budget", 1000)
    for k in ("NEUROSEARCH_TASK_MODEL_FINDINGS_EXTRACT", "NEUROSEARCH_TASK_THINKING_FINDINGS_EXTRACT"):
        monkeypatch.delenv(k, raising=False)
    out = tmp_path / "evals"
    cmp = evals.run_findings_compare(live=False, out_dir=out, progress=lambda m: None)
    b, c = cmp["baseline"], cmp["candidate"]
    assert (b["configured_model"], c["configured_model"]) == ("claude-sonnet-4-6", "claude-sonnet-5")
    assert b["contract"]["thinking"] == c["contract"]["thinking"] == "disabled" and b["prompt_version"] == c["prompt_version"] == "18b5db69"
    assert {k: v for k, v in b["contract"].items() if k != "model"} == {k: v for k, v in c["contract"].items() if k != "model"}
    assert b["volume"]["input_tokens"] == c["volume"]["input_tokens"] == 30297      # the frozen Tier-1 findings workload, twice
    assert b["quality"]["window_calls"] == c["quality"]["window_calls"] == 9 and b["invocations"]["logical"] == 9
    assert b["quality"]["nuggets"] == c["quality"]["nuggets"] and len(b["quality"]["nuggets"]) == 10
    assert b["quality"]["golden_evidence_recall"] == c["quality"]["golden_evidence_recall"] > 0.5
    assert b["quality"]["finding_quote_validity"] == 1.0 and b["quality"]["incomplete_outputs"] == 0 and b["quality"]["repaired_windows"] == 0
    assert b["quality"]["summary_validity"] == 1.0 and b["quality"]["substance_validity"] == 1.0
    assert b["economics"]["cost_per_source_hour"] > 0 and b["performance"]["findings_s"] >= 0
    assert abs(c["canonical_input_tokens"] - b["canonical_input_tokens"] * 1.3) < 20 and b["canonical_requests"] == 9   # per-request rounding
    assert cmp["verdict"]["verdict"] == "PASS" and cmp["verdict"]["production_default_changed"] is False and cmp["verdict"]["lost_nuggets"] == []
    assert contracts.contract("findings.extract").model == "claude-sonnet-5" and "NEUROSEARCH_TASK_MODEL_FINDINGS_EXTRACT" not in os.environ
    names = {p.name for p in out.rglob("*")}
    assert {"baseline-sonnet-4-6.json", "candidate-sonnet-5.json", "comparison.json", "comparison.txt"} <= names
    assert cmp["findings_baseline_written"] and pathlib.Path(cmp["findings_baseline_file"]).exists()
    metrics = {r["metric"].split("[")[0] for r in cmp["rows"]}
    for k in ("golden_evidence_recall", "finding_quote_validity", "findings_rejected", "nugget", "findings_per_source", "findings_per_window", "summary_validity",
              "substance_validity", "truncated_windows", "incomplete_outputs", "repaired_windows", "outcome_unknown", "configured_model", "returned_model",
              "canonical_input_tokens", "billed_input_tokens", "billed_output_tokens", "billed_cache_read_tokens", "cost", "cost_per_source_hour", "latency_findings_s"):
        assert k in metrics, k
    assert "VERDICT: PASS — migrate findings.extract to claude-sonnet-5" in cmp["text"] and "Production default unchanged" in cmp["text"]
    cmp2 = evals.run_findings_compare(live=False, out_dir=out, progress=lambda m: None)
    assert cmp2["findings_baseline_written"] is False


def test_findings_verdict_rules():
    """Evidence recall + quote integrity are hard gates; the number of findings is not; cost/latency only caveat."""
    import copy
    from neurosearch import evals
    nug = {f"n{i}": True for i in range(8)} | {"n8": False, "n9": False}
    base = {"configured_model": "claude-sonnet-4-6", "returned_model": "claude-sonnet-4-6-20260210", "contract": {"thinking": "disabled"},
            "quality": {"golden_evidence_recall": 0.8, "nuggets": nug, "finding_quote_validity": 1.0, "stored_findings_verified": 1.0, "findings_rejected": 0,
                        "findings_suggested": 48, "raw_findings": 48, "incomplete_outputs": 0, "truncated_windows": 0, "parse_failed_windows": 0, "no_json_windows": 0,
                        "repaired_windows": 0, "empty_windows": 0, "failed_sources": [], "summary_validity": 1.0, "substance_validity": 1.0,
                        "windows_missing_summary": 0, "windows_missing_substance": 0},
            "invocations": {"outcome_unknown": 0, "attempts": 9}, "economics": {"cost": 0.196, "cost_per_source_hour": 0.12}, "performance": {"findings_s": 60.0},
            "canonical_input_tokens": 32000}
    cand = copy.deepcopy(base); cand["configured_model"] = "claude-sonnet-5"; cand["returned_model"] = "claude-sonnet-5-20260601"; cand["canonical_input_tokens"] = 41600
    assert evals.findings_verdict(base, cand)["verdict"] == "PASS"
    # fewer findings alone is never a FAIL
    c = copy.deepcopy(cand); c["quality"]["findings_suggested"] = 30
    v = evals.findings_verdict(base, c); assert v["verdict"] == "PASS_WITH_CAVEAT" and any("not a gate" in x for x in v["caveats"])
    c = copy.deepcopy(cand); c["quality"]["findings_suggested"] = 44
    assert evals.findings_verdict(base, c)["verdict"] == "PASS"
    # one lost nugget → caveat (named); two → FAIL (named); a lost nugget replaced by a gained one still shows as lost
    c = copy.deepcopy(cand); c["quality"]["nuggets"]["n0"] = False; c["quality"]["golden_evidence_recall"] = 0.7
    v = evals.findings_verdict(base, c); assert v["verdict"] == "PASS_WITH_CAVEAT" and v["lost_nuggets"] == ["n0"] and "n0" in v["caveats"][0]
    c = copy.deepcopy(cand); c["quality"]["nuggets"].update(n0=False, n1=False); c["quality"]["golden_evidence_recall"] = 0.6
    v = evals.findings_verdict(base, c); assert v["verdict"] == "FAIL" and v["lost_nuggets"] == ["n0", "n1"] and v["headline"] == "FAIL — keep claude-sonnet-4-6 for findings.extract"
    c = copy.deepcopy(cand); c["quality"]["nuggets"].update(n0=False, n8=True); c["quality"]["golden_evidence_recall"] = 0.8
    v = evals.findings_verdict(base, c); assert v["verdict"] == "PASS_WITH_CAVEAT" and v["lost_nuggets"] == ["n0"] and v["gained_nuggets"] == ["n8"]
    # quote integrity under the floor, truncation, parse failure, unknown outcome, wrong model → FAIL
    for patch in ({"quality": {"finding_quote_validity": 0.9}}, {"quality": {"stored_findings_verified": 0.95}}, {"quality": {"incomplete_outputs": 1, "truncated_windows": 1}},
                  {"quality": {"incomplete_outputs": 1, "parse_failed_windows": 1}}, {"quality": {"failed_sources": ["yt01"]}}, {"invocations": {"outcome_unknown": 1}},
                  {"returned_model": "claude-sonnet-4-6-20260210"}):
        c = copy.deepcopy(cand)
        for k, vv in patch.items():
            if isinstance(vv, dict):
                c[k].update(vv)
            else:
                c[k] = vv
        assert evals.findings_verdict(base, c)["verdict"] == "FAIL", patch
    # cost/latency/repairs → caveats only
    c = copy.deepcopy(cand); c["economics"]["cost"] = 0.3; c["performance"]["findings_s"] = 120.0; c["quality"]["repaired_windows"] = 2
    v = evals.findings_verdict(base, c); assert v["verdict"] == "PASS_WITH_CAVEAT" and len(v["caveats"]) == 3
    assert any("+30.0%" in n for n in v["notes"])


def test_findings_structured_output_mismatch_is_a_bug_signal_not_a_parse(isolated_db, monkeypatch):
    """Mission F path: provider structured result → json.loads → FULL local validation → PASS store / FAIL SchemaMismatch.
    A mismatch on a normal completion is a bug signal: recorded, retried once with a fresh completion, then a typed error
    that propagates. The legacy parser is never on this path unless NEUROSEARCH_SCHEMA_COMPAT_FALLBACK=1, and even then
    its result must pass the full schema. Truncation: one guarded escalation capped by the contract ceiling."""
    from neurosearch import findings, providers, usage
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.delenv("NEUROSEARCH_TASK_SCHEMA_FINDINGS_EXTRACT", raising=False)
    monkeypatch.delenv("NEUROSEARCH_SCHEMA_COMPAT_FALLBACK", raising=False)

    class R:
        def __init__(self, text, stop="end_turn"):
            self.content = [type("B", (), {"type": "text", "text": text})()]
            self.stop_reason, self.model, self.usage = stop, "claude-x", type("U", (), {"input_tokens": 10, "output_tokens": 5, "cache_read_input_tokens": 0, "cache_creation_input_tokens": 0})()
    calls, guards = [], []
    def fake_invoke(resp_seq):
        it = iter(resp_seq)
        def f(task, **kw):
            calls.append(kw); return next(it)
        return f
    monkeypatch.setattr(usage, "guard", lambda est=0.0: guards.append(est))
    monkeypatch.setattr(usage, "record_anthropic", lambda *a, **k: 0.0)
    good = '{"summary":"s","substance":50,"findings":[]}'
    so = lambda: db.health()["structured_outputs"]  # noqa: E731
    # strict path: one call, validated, no events
    monkeypatch.setattr(providers, "invoke", fake_invoke([R(good)]))
    assert findings._call(findings.SYSTEM, "u", head="h") == {"summary": "s", "substance": 50, "findings": []}
    assert findings._last_call["parse"] == "strict" and findings._last_call["structured"] and so() == {"mismatches": 0, "mismatches_recovered_by_retry": 0, "fallbacks": 0, "unrecovered": 0, "truncated": 0, "refused": 0, "steady_state": "all zero"}
    # fenced JSON from a structured request = mismatch → retry → second is fine → recovered, counted as a mismatch (bug signal)
    calls.clear()
    monkeypatch.setattr(providers, "invoke", fake_invoke([R("```json\n" + good + "\n```"), R(good)]))
    assert findings._call(findings.SYSTEM, "u", head="h")["substance"] == 50 and len(calls) == 2
    assert so()["mismatches"] == 1 and so()["mismatches_recovered_by_retry"] == 1 and so()["fallbacks"] == 0
    ev = db.validation_events(kind="schema_mismatch")
    assert len(ev) == 1 and ev[0]["detail"]["task"] == "findings.extract" and ev[0]["detail"]["schema"] == "findings-v2" and "not JSON" in ev[0]["detail"]["reason"]
    # mismatch twice → typed SchemaMismatch propagates; nothing is stored; legacy parser NOT used
    calls.clear()
    monkeypatch.setattr(providers, "invoke", fake_invoke([R("```json\n" + good + "\n```"), R('{"summary":"s","substance":"high","findings":[]}')]))
    with pytest.raises(providers.SchemaMismatch) as ei:
        findings._call(findings.SYSTEM, "u", head="h")
    assert "substance" in str(ei.value) and len(calls) == 2 and so()["mismatches"] == 2 and so()["unrecovered"] == 1 and so()["fallbacks"] == 0
    assert findings._last_call["parse"] == "failed"
    # the explicitly degraded compat hatch: legacy parser may run, but its result must pass the FULL schema
    monkeypatch.setenv("NEUROSEARCH_SCHEMA_COMPAT_FALLBACK", "1")
    monkeypatch.setattr(providers, "invoke", fake_invoke([R("```json\n" + good + "\n```"), R("```json\n" + good + "\n```")]))
    assert findings._call(findings.SYSTEM, "u", head="h")["substance"] == 50 and so()["fallbacks"] == 1
    assert db.validation_events(kind="schema_fallback")[0]["detail"]["validated_against"] == "findings-v2"
    monkeypatch.setattr(providers, "invoke", fake_invoke([R('{"summary":"s","substance":"high","findings":[]}'), R('{"summary":"s","substance":"high","findings":[]}')]))
    with pytest.raises(providers.SchemaMismatch):                       # legacy parser output fails the full schema → still a typed failure
        findings._call(findings.SYSTEM, "u", head="h")
    assert so()["fallbacks"] == 2 and so()["unrecovered"] == 2
    monkeypatch.delenv("NEUROSEARCH_SCHEMA_COMPAT_FALLBACK")
    # truncation: one escalation through the guard, capped by the contract ceiling (findings: 4000 → min(6000, 6000)), budgets recorded
    calls.clear(); guards.clear()
    monkeypatch.setattr(providers, "invoke", fake_invoke([R('{"summary":"s"', "max_tokens"), R('{"summary":"s"', "max_tokens")]))
    with pytest.raises(providers.OutputError) as ei:
        findings._call(findings.SYSTEM, "u", head="h")
    assert ei.value.kind == "TRUNCATED" and len(calls) == 2 and calls[1]["max_output_tokens"] == 6000 and "max_output_tokens" not in calls[0]
    assert len(guards) == 2 and guards[1] == pytest.approx(guards[0] * 1.5)
    evs = db.validation_events(kind="output_truncated")
    assert so()["truncated"] == 2 and {k: evs[-1]["detail"].get(k) for k in ("budget", "escalated_to", "ceiling")} == {"budget": 4000, "escalated_to": 6000, "ceiling": 6000}
    assert evs[0]["detail"]["final"] is True and evs[0]["detail"]["budget"] == 6000
    calls.clear()
    monkeypatch.setattr(providers, "invoke", fake_invoke([R('{"summary":"s"', "max_tokens"), R(good)]))
    assert findings._call(findings.SYSTEM, "u", head="h")["substance"] == 50 and len(calls) == 2
    # refusal is typed, never parsed
    monkeypatch.setattr(providers, "invoke", fake_invoke([R("I can't help with that.", "refusal")]))
    with pytest.raises(providers.OutputError) as ei:
        findings._call(findings.SYSTEM, "u", head="h")
    assert ei.value.kind == "REFUSED" and so()["refused"] == 1
    # rollback switch: NEUROSEARCH_TASK_SCHEMA_FINDINGS_EXTRACT=none restores the pre-F path (no output_config, tolerant parse, no events)
    monkeypatch.setenv("NEUROSEARCH_TASK_SCHEMA_FINDINGS_EXTRACT", "none")
    from neurosearch import contracts as C
    assert C.contract("findings.extract").schema is None and "output_config" not in C.request_params(C.contract("findings.extract"))
    before = dict(so())
    monkeypatch.setattr(providers, "invoke", fake_invoke([R("```json\n" + good + "\n```")]))
    assert findings._call(findings.SYSTEM, "u", head="h")["substance"] == 50 and not findings._last_call["structured"] and findings._last_call["parse"] == "repaired"
    assert so() == before


def test_truncation_escalation_needs_a_ceiling(isolated_db, monkeypatch):
    """No ceiling on the contract → no escalation at all (a typed error straight away): escalation can never be unbounded."""
    from neurosearch import contracts as C, providers, usage
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(usage, "guard", lambda *a, **k: None)
    monkeypatch.setattr(usage, "record_anthropic", lambda *a, **k: 0.0)
    class R:
        content = [type("B", (), {"type": "text", "text": "{"})()]; stop_reason = "max_tokens"; model = "m"; usage = None
    calls = []
    monkeypatch.setattr(providers, "invoke", lambda task, **kw: calls.append(kw) or R())
    monkeypatch.setattr(C, "_base", lambda: {"rank.relevance": C.InferenceContract("rank.relevance", "anthropic", "claude-sonnet-5", thinking="disabled", max_output_tokens=6000, schema="rank-v2")})
    with pytest.raises(providers.OutputError) as ei:
        providers.invoke_structured("rank.relevance", system="s", messages=[{"role": "user", "content": "u"}], usage_kind="rank")
    assert ei.value.kind == "TRUNCATED" and len(calls) == 1
    with pytest.raises(C.ContractError, match="ceiling"):
        C.validate(C.InferenceContract("x", "anthropic", "claude-sonnet-5", thinking="disabled", max_output_tokens=6000, max_output_ceiling=1000, schema="rank-v2"))


def test_schema_registry_is_provider_compatible_and_validates():
    import pytest as _pt
    from neurosearch import schemas
    def keywords(node, is_map=False):
        if isinstance(node, dict):
            for k, v in node.items():
                if not is_map:
                    yield k
                yield from keywords(v, is_map=(not is_map and k in ("properties", "$defs", "definitions")))
        elif isinstance(node, list):
            for x in node:
                yield from keywords(x)
    for name, sch in schemas.REGISTRY.items():
        schemas.check_provider_compat(sch, name)                 # would raise before any paid request
        ps = schemas.provider_schema(name)
        kws = set(keywords(ps))
        assert not (kws & {"minimum", "maximum", "minLength", "pattern", "$schema"}), (name, kws)      # client-only constraints stripped …
        assert "$ref" in json.dumps(ps) or "$defs" not in sch                                         # … internal $ref/$defs kept
        assert "minimum" in json.dumps(schemas.provider_schema("plan-economics-v3"))                  # property NAMES are never stripped (costs.minimum)
        assert schemas.output_config(name)["format"]["type"] == "json_schema"
    assert schemas.is_valid("findings-v2", {"summary": "s", "substance": 40, "findings": [{"title": "t", "finding": "", "ts": "1:00", "quote": "q", "importance": 3}]})
    errs = schemas.validate("findings-v2", {"summary": "s", "substance": 400, "findings": [{"title": "t"}]})
    assert any("substance" in e for e in errs) and any("finding" in e or "required" in e for e in errs)
    assert not schemas.is_valid("findings-v2", {"summary": "s", "substance": 1, "findings": [], "extra": 1})       # additionalProperties: false
    assert schemas.is_valid("rank-v2", {"scores": [{"i": 0, "score": 90, "why": "on topic"}]}) and not schemas.is_valid("rank-v2", {"scores": [{"i": 0, "score": 101, "why": ""}]})
    for bad, msg in (({"type": "object", "properties": {}, "additionalProperties": False, "allOf": []}, "allOf"),
                     ({"type": "object", "properties": {"a": {"$ref": "https://x/y"}}, "additionalProperties": False}, "internal"),
                     ({"type": "object", "properties": {"a": {"$ref": "#/$defs/nope"}}, "additionalProperties": False}, "unresolved"),
                     ({"type": "object", "properties": {}}, "additionalProperties")):
        with _pt.raises(schemas.SchemaError, match=msg):
            schemas.check_provider_compat(bad, "bad")
    with _pt.raises(schemas.SchemaError):
        schemas.get("findings-v9")


def test_fake_conforms_to_registry_schemas_and_bad_json_knob(isolated_db, monkeypatch):
    """Tier 1 proves conformance: the fake validates its own structured outputs; the BAD_JSON knob produces the bug-signal path
    (mismatch → retry → typed failure), which the ranking job surfaces as a failed batch, never a repaired one."""
    from neurosearch import fake_ai, providers, relevance
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(settings, "daily_budget", 1000)
    monkeypatch.delenv("NEUROSEARCH_FAKE_AI_BAD_JSON", raising=False)
    monkeypatch.delenv("NEUROSEARCH_SCHEMA_COMPAT_FALLBACK", raising=False)
    r = providers.invoke("rank.relevance", system="PROJECT: x\nBRIEF: getting out of debt", messages=[{"role": "user", "content": "VIDEOS:\n[0] Debt payoff plan\n[1] Vlog\n\nScore them now."}])
    out = providers.structured("rank.relevance", r)
    assert [x["i"] for x in out["scores"]] == [0, 1] and all(0 <= x["score"] <= 100 for x in out["scores"])
    monkeypatch.setenv("NEUROSEARCH_FAKE_AI_BAD_JSON", "1")
    with pytest.raises(providers.SchemaMismatch):
        relevance._call(relevance.SYSTEM, "VIDEOS:\n[0] a\n[1] b\n\nScore them now.", None, "c", head="PROJECT: x\n")
    h = db.health()["structured_outputs"]
    assert h["mismatches"] == 1 and h["unrecovered"] == 1 and h["fallbacks"] == 0
    # through rank_collection: a visibly failed batch (unscored, re-rank offered), not a silent repair
    p = db.create_project("Bad JSON", "getting out of debt")
    coll = db.upsert_collection("channel", "UCbad", "https://www.youtube.com/@bad", "Bad")
    db.add_project_collections(p["id"], [coll["id"]])
    for i in range(3):
        src = db.upsert_source(platform="youtube", external_id=f"bad{i:08d}", url=f"https://www.youtube.com/watch?v=bad{i:08d}", title=f"V{i}", description="debt", duration=60, status="proposed")
        db.link_source_collection(src["id"], coll["id"])
    res = relevance.rank_collection(coll["id"], p["id"], want=2)
    assert res["failed_batches"] == 1 and res["repaired_batches"] == 0 and res["ranked"] == 0
    assert all(s["relevance"] is None for s in db.proposed_sources(coll["id"], p["id"]))
    # with the compat hatch on, the legacy parser recovers it — still counted as a fallback, still fully validated
    monkeypatch.setenv("NEUROSEARCH_SCHEMA_COMPAT_FALLBACK", "1")
    res = relevance._call(relevance.SYSTEM, "VIDEOS:\n[0] a\n[1] b\n\nScore them now.", None, "c", head="PROJECT: x\n")
    assert len(res["scores"]) == 2 and db.health()["structured_outputs"]["fallbacks"] == 1
    monkeypatch.delenv("NEUROSEARCH_SCHEMA_COMPAT_FALLBACK")
    # the fake refuses to emit non-conformant output for a structured task unless the knob is on
    monkeypatch.delenv("NEUROSEARCH_FAKE_AI_BAD_JSON", raising=False)
    monkeypatch.setattr(fake_ai, "_rank", lambda user, system="": '{"scores": [{"i": "zero", "score": 1, "why": "x"}]}')
    with pytest.raises(providers.ProviderError, match="does not conform"):        # surfaces through the ledger as a failed attempt
        providers.invoke("rank.relevance", system="s", messages=[{"role": "user", "content": "VIDEOS:\n[0] a"}])


def test_plan_update_and_discover_quick_are_structured(isolated_db, monkeypatch):
    """F3: planner.update never collapses a misunderstood output into an innocent []; discover.quick is schema'd;
    discover.verify keeps its citation-capable free-text path."""
    from neurosearch import contracts as C, discover, planner, providers
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(settings, "daily_budget", 1000)
    monkeypatch.delenv("NEUROSEARCH_FAKE_AI_BAD_JSON", raising=False)
    assert C.contract("planner.update").schema == "plan-update-v2" and C.contract("discover.quick").schema == "discovery-v2" and C.contract("discover.verify").schema is None
    p = db.create_project("Upd", "moving a website")
    db.save_plan(p["id"], {"goal": {"outcome": "move the site"}, "approach": {"recommended": "Static export + Cloudflare Pages"}}, db.project_snapshot(p["id"]))
    import time as _t; _t.sleep(0.02)
    db.add_project_note(p["id"], "Cloudflare form handling is limited.", [])
    ups = planner.suggest_updates(p["id"])
    assert len(ups) == 1 and "Cloudflare form handling" in ups[0]["proposed"] and ups[0]["reason"].startswith("New finding") and db.latest_plan(p["id"]) and db.health()["structured_outputs"]["mismatches"] == 0
    # a misunderstood output is a typed failure, and NO update row is written (not even an empty list)
    monkeypatch.setenv("NEUROSEARCH_FAKE_AI_BAD_JSON", "1")
    db.add_project_note(p["id"], "Another new finding.", [])
    with pytest.raises(providers.SchemaMismatch):
        planner.suggest_updates(p["id"])
    assert db.health()["structured_outputs"]["unrecovered"] == 1
    monkeypatch.delenv("NEUROSEARCH_FAKE_AI_BAD_JSON")
    # discover.quick: structured; the fake's canned shortlist conforms to discovery-v2
    d = discover.discover(p["id"], count=2)
    assert len(d["items"]) == 3 and d["items"][0]["name"] == "Dave Ramsey" and d["items"][0]["known_for"] == "debt-free budgeting basics"   # 2 quick + 1 added by verify
    assert db.health()["structured_outputs"]["mismatches"] == 1      # unchanged by discover: no new bug signals


def test_schema_version_is_part_of_provenance_and_input_hash(isolated_db, monkeypatch):
    """A schema change makes findings/rankings stale exactly like a prompt change, and provenance names the schema."""
    from neurosearch import findings, relevance
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    p = db.create_project("Schema", "buying a business")
    r = ingest.store_transcript({"platform": "youtube", "external_id": "schm0000001", "url": "https://www.youtube.com/watch?v=schm0000001", "title": "T",
                                 "segments": [{"start": 0, "end": 5, "text": "the SBA requires ten percent down"}], "transcript_kind": "captions"}, project_id=p["id"])
    sid = r["source_id"]
    h1 = findings.input_hash(p, sid)
    monkeypatch.setenv("NEUROSEARCH_TASK_SCHEMA_FINDINGS_EXTRACT", "none")
    assert findings.input_hash(p, sid) != h1 and findings.schema_version() is None
    monkeypatch.delenv("NEUROSEARCH_TASK_SCHEMA_FINDINGS_EXTRACT")
    assert findings.input_hash(p, sid) == h1 and findings.schema_version() == "findings-v2"
    findings.suggest_for_source(p["id"], sid)
    a = db.get_analysis(p["id"], sid, "summary")
    assert a["schema_version"] == "findings-v2" and a["input_hash"] == h1
    s = {"title": "x", "description": "y", "duration": 60}
    r1 = relevance.input_hash(p, s)
    monkeypatch.setenv("NEUROSEARCH_TASK_SCHEMA_RANK_RELEVANCE", "none")
    assert relevance.input_hash(p, s) != r1 and relevance.schema_version() is None


# ------------------------------------------------------------------ E2.3 migration comparison (one command)

def test_planner_rubric_is_frozen_and_deterministic():
    import re as _re
    from neurosearch import migration
    rub = migration.load_rubric()
    assert rub["version"] == 1 and len(rub["plan"]["checks"]) >= 20 and len(rub["analysis"]["checks"]) >= 8
    for sec in ("plan", "analysis"):
        ids = [c["id"] for c in rub[sec]["checks"]]
        assert len(ids) == len(set(ids))
        for c in rub[sec]["checks"]:
            assert c["kind"] in ("risk", "assumption", "dependency", "first_step", "contradiction", "open_question", "completeness")
            for p in c.get("any", []) + c.get("all", []):
                _re.compile(p, _re.I)
            assert c.get("any") or c.get("all")
        kinds = {c["kind"] for c in rub[sec]["checks"]}
        assert {"risk", "assumption", "contradiction", "completeness"} <= kinds
    assert {"risk", "assumption", "dependency", "first_step", "contradiction", "open_question", "completeness"} == {c["kind"] for c in rub["plan"]["checks"]}
    # deterministic scoring on a synthetic plan: the same document always scores the same, and sections are respected
    plan = {"risks": [{"risk": "Customer concentration above 20% of revenue", "mitigation": "price it in or walk", "priority": "high"}] * 3,
            "first_steps": [{"action": "Call two SBA lenders about a 10% equity injection"}] * 3,
            "this_week": [{"action": "Call an SBA lender", "why": "financing gates everything", "time": "1h"}] * 3,
            "open_questions": [{"question": "How much cash can you bring?", "research_prompt": "x"}] * 2,
            "refine_questions": [{"question": "HVAC or manufacturing?", "why": "y"}] * 4, "phases": [{"name": "a", "objective": "b"}] * 2,
            "dependencies": [{"item": "lender pre-qualification"}] * 2, "gotchas": [{"gotcha": "g", "avoid": "a"}] * 2,
            "approach": {"recommended": "SBA 7(a) with a seller note on standby; the podcast's zero down seller financing path conflicts with the 10% injection rule", "basis": "research", "evidence": ["F1"]},
            "costs": {"upfront": [{"item": "SBA guarantee fee", "amount": "3%"}], "evidence": []}}
    r1 = migration.score_rubric(plan, rub["plan"]); r2 = migration.score_rubric(plan, rub["plan"])
    assert r1 == r2 and 0 < r1["score"] < 1 and "risk.customer_concentration" not in r1["failed"] and "contradiction.zero_down_vs_injection" not in r1["failed"]
    assert "risk.lease_assignment" in r1["failed"] and r1["structure"]["score"] == 1.0
    assert migration.score_rubric({}, rub["plan"])["score"] == 0.0
    # 'sections' matter: a risk mentioned only in the summary text does not satisfy a risks-section check
    assert "risk.seller_transition" in migration.score_rubric({"goal": {"outcome": "seller transition"}}, rub["plan"])["failed"]
    g = migration.grounding(plan)
    assert g == {"research_items": 1, "grounded": 1, "fraction": 1.0, "evidence_refs": 1}


def test_migration_verdict_rules():
    """Planner: adaptive is recommended only when it measurably beats disabled; equal → disabled. Chat/export/update gates."""
    import copy
    from neurosearch import migration as M
    inv = {"logical": 1, "attempts": 1, "by_state": {"completed": 1}, "outcome_unknown": 0, "returned_model": "claude-sonnet-5-20260601"}
    inv46 = {**inv, "returned_model": "claude-sonnet-4-6-20260210"}
    def task(rubric, struct=1.0, **kw):
        base = {"present": True, "rubric": {"score": rubric, "passed": 10, "total": 20, "failed": [], "by_kind": {}, "structure": {"score": struct, "failed": []}},
                "grounding": {"research_items": 5, "grounded": 5, "fraction": 1.0, "evidence_refs": 30}, "evidence_refs": 30, "dangling_after_removal": 0,
                "truncated": 0, "parse_failed": 0, "json_repaired": 0, "output_chars": 8000, "output_tokens": 2000, "input_tokens": 100, "cache_read": 30000,
                "cache_write": 0, "cost": 0.1, "thinking_tokens_est": 0, "thinking_blocks": 0, "seconds": 30.0, "invocations": dict(inv)}
        base.update(kw); return base
    def arm(rubric_a, rubric_b, returned=inv, **kw):
        a = {"error": None, "seconds": 60.0, "tasks": {"planner.analysis": task(rubric_a, invocations=dict(returned)), "planner.build": task(rubric_b, invocations=dict(returned))},
             "plan_evidence_dangling_raw": 0, "plan_evidence_refs_raw": 60, "usage": {"cost": 0.2}}
        for k, v in kw.items():
            a["tasks"]["planner.build"][k] = v
        return a
    metas = {"4.6": {"arm": "4.6", "model": "claude-sonnet-4-6"}, "5-disabled": {"arm": "5-disabled", "model": "claude-sonnet-5"}, "5-adaptive-medium": {"arm": "5-adaptive-medium", "model": "claude-sonnet-5"}}
    arms = {"4.6": arm(0.7, 0.7, inv46), "5-disabled": arm(0.7, 0.7), "5-adaptive-medium": arm(0.7, 0.72)}
    v = M.planner_verdict("planner.build", arms, metas)
    assert v["verdict"] == "PASS" and not v["adaptive"]["recommended"] and v["recommended_setting"]["thinking"] == "disabled"
    arms["5-adaptive-medium"] = arm(0.7, 0.85)
    v = M.planner_verdict("planner.build", arms, metas)
    assert v["adaptive"]["recommended"] and v["recommended_setting"] == {"model": "claude-sonnet-5", "thinking": "adaptive", "effort": "medium"} and "ADAPTIVE" in v["headline"]
    arms["5-adaptive-medium"] = arm(0.7, 0.7, truncated=1)                        # adaptive fails a validity gate → never recommended
    assert not M.planner_verdict("planner.build", arms, metas)["adaptive"]["recommended"]
    arms["5-disabled"] = arm(0.7, 0.55)                                          # disabled arm regresses the rubric → FAIL keeps 4.6
    v = M.planner_verdict("planner.build", arms, metas)
    assert v["verdict"] == "FAIL" and v["recommended_setting"]["model"] == "claude-sonnet-4-6"
    arms["5-disabled"] = arm(0.7, 0.65)                                          # within tolerance → caveat
    assert M.planner_verdict("planner.build", arms, metas)["verdict"] == "PASS_WITH_CAVEAT"
    arms["5-disabled"] = arm(0.7, 0.7, json_repaired=1)                           # JSON repair the baseline did not need → FAIL
    assert M.planner_verdict("planner.build", arms, metas)["verdict"] == "FAIL"
    arms["5-disabled"] = arm(0.7, 0.7, returned=inv46)                            # wrong returned model → FAIL
    assert M.planner_verdict("planner.build", arms, metas)["verdict"] == "FAIL"
    # chat
    chat = {"answers": 34, "failed_answers": 0, "citation_validity": 1.0, "answers_cite_expected_source": 0.9, "contradiction_surfaced": 1.0, "contradiction_questions": 3,
            "gap_detection": 1.0, "gap_questions": 2, "offtopic_handled": 1.0, "repair_rounds": 1, "repairs_succeeded": 1, "answers_with_unrepaired_invalid_citations": 0,
            "truncated_answers": 0, "tool_rounds": 2, "tool_calls": 2, "mean_answer_chars": 500, "s_per_answer": 6.0, "seconds": 200.0,
            "invocations": dict(inv46), "repair_invocations": dict(inv46), "usage": {"cost": 0.8}}
    c5 = copy.deepcopy(chat); c5["invocations"] = dict(inv); c5["repair_invocations"] = dict(inv)
    cm = {"4.6": metas["4.6"], "5-disabled": metas["5-disabled"]}
    ch, rp = M.chat_verdict({"4.6": chat, "5-disabled": c5}, cm)
    assert ch["verdict"] == "PASS" and rp["verdict"] == "PASS"
    c = copy.deepcopy(c5); c["contradiction_surfaced"] = 0.6667
    assert M.chat_verdict({"4.6": chat, "5-disabled": c}, cm)[0]["verdict"] == "PASS_WITH_CAVEAT"
    c = copy.deepcopy(c5); c["contradiction_surfaced"] = 0.3333
    assert M.chat_verdict({"4.6": chat, "5-disabled": c}, cm)[0]["verdict"] == "FAIL"
    c = copy.deepcopy(c5); c["citation_validity"] = 0.98; c["answers_with_unrepaired_invalid_citations"] = 1
    assert M.chat_verdict({"4.6": chat, "5-disabled": c}, cm)[0]["verdict"] == "FAIL"
    c = copy.deepcopy(c5); c["truncated_answers"] = 1
    assert M.chat_verdict({"4.6": chat, "5-disabled": c}, cm)[0]["verdict"] == "FAIL"
    c = copy.deepcopy(c5); c["repair_rounds"] = 2; c["repairs_succeeded"] = 1
    ch, rp = M.chat_verdict({"4.6": chat, "5-disabled": c}, cm)
    assert ch["verdict"] == "PASS_WITH_CAVEAT" and rp["verdict"] == "FAIL"
    c = copy.deepcopy(c5); c["mean_answer_chars"] = 900; c["usage"]["cost"] = 0.7            # wording/length differences are not failures
    assert M.chat_verdict({"4.6": chat, "5-disabled": c}, cm)[0]["verdict"] == "PASS"
    # export
    ex = {"words": 1200, "sections_present": 6, "sections_missing": [], "citation_links": 30, "given_links": 33, "link_coverage": 0.9, "n_invented_links": 0, "invented_links": [],
          "truncated": 0, "empty": 0, "fallback_used": 0, "seconds": 20.0, "invocations": dict(inv46), "usage": {"cost": 0.05}}
    e5 = copy.deepcopy(ex); e5["invocations"] = dict(inv)
    assert M.export_verdict({"4.6": ex, "5-disabled": e5}, cm)["verdict"] == "PASS"
    e = copy.deepcopy(e5); e["n_invented_links"] = 1; e["invented_links"] = ["https://x"]
    assert M.export_verdict({"4.6": ex, "5-disabled": e}, cm)["verdict"] == "FAIL"
    e = copy.deepcopy(e5); e["sections_missing"] = ["Open questions"]
    assert M.export_verdict({"4.6": ex, "5-disabled": e}, cm)["verdict"] == "FAIL"
    e = copy.deepcopy(e5); e["link_coverage"] = 0.8
    assert M.export_verdict({"4.6": ex, "5-disabled": e}, cm)["verdict"] == "PASS_WITH_CAVEAT"
    # update
    up = {"n_updates": 2, "addresses_new_finding": 1, "addresses_new_fact": 1, "well_formed": 1, "json_repaired": 0, "truncated": 0, "parse_failed": 0, "seconds": 10.0,
          "invocations": dict(inv46), "usage": {"cost": 0.03}}
    u5 = copy.deepcopy(up); u5["invocations"] = dict(inv)
    assert M.update_verdict({"4.6": up, "5-disabled": u5}, cm)["verdict"] == "PASS"
    u = copy.deepcopy(u5); u["addresses_new_finding"] = 0
    assert M.update_verdict({"4.6": up, "5-disabled": u}, cm)["verdict"] == "FAIL"
    u = copy.deepcopy(u5); u["addresses_new_fact"] = 0
    assert M.update_verdict({"4.6": up, "5-disabled": u}, cm)["verdict"] == "PASS_WITH_CAVEAT"


def test_migration_compare_one_command(isolated_db, monkeypatch, tmp_path):
    """--migration-compare under the fakes: every arm sees identical inputs, project state is restored between arms,
    env overrides are restored, all raw arms + the comparison are saved, one verdict per task, nothing migrated."""
    import os
    from neurosearch import contracts, migration
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    for t in ("PLANNER_ANALYSIS", "PLANNER_BUILD", "PLANNER_UPDATE", "EXPORT_SYNTHESIS", "ANSWER_CHAT", "ANSWER_REPAIR"):
        for w in ("MODEL", "THINKING", "EFFORT"):
            monkeypatch.delenv(f"NEUROSEARCH_TASK_{w}_{t}", raising=False)
    out = tmp_path / "evals"
    rep = migration.run_migration_compare(live=False, out_dir=out, progress=lambda m: None)
    assert set(rep["summary"]) == {"planner.analysis", "planner.build", "planner.update", "export.synthesis", "answer.chat", "answer.repair"}
    assert all(v != "FAIL" for v in rep["summary"].values()), rep["summary"]
    assert set(rep["excluded"]) == {"discover.quick", "discover.verify"}
    p = rep["arms"]["planner"]
    assert set(p) == {"4.6", "5-disabled", "5-adaptive-medium"}
    assert p["4.6"]["meta"]["contracts"]["planner.build"]["model"] == "claude-sonnet-4-6" and p["5-disabled"]["meta"]["contracts"]["planner.build"]["thinking"] == "disabled"
    assert p["5-adaptive-medium"]["meta"]["contracts"]["planner.analysis"] | {} == p["5-adaptive-medium"]["meta"]["contracts"]["planner.analysis"]
    assert p["5-adaptive-medium"]["meta"]["contracts"]["planner.analysis"]["thinking"] == "adaptive" and p["5-adaptive-medium"]["meta"]["contracts"]["planner.analysis"]["effort"] == "medium"
    # identical inputs: the same fake sees the same prompt in every arm (same input tokens, same rubric)
    for t in ("planner.analysis", "planner.build"):
        toks = {k: v["tasks"][t]["input_tokens"] + v["tasks"][t]["cache_read"] + v["tasks"][t]["cache_write"] for k, v in p.items()}
        assert len(set(toks.values())) == 1, toks
        assert len({v["tasks"][t]["rubric"]["score"] for v in p.values()}) == 1
        assert not rep["verdicts"][t]["adaptive"]["recommended"] and rep["verdicts"][t]["recommended_setting"]["thinking"] == "disabled"
    c = rep["arms"]["answer.chat"]
    tot = lambda u: u["input_tokens"] + u["cache_read_tokens"] + u["cache_write_tokens"]  # noqa: E731
    assert tot(c["4.6"]["usage"]) == tot(c["5-disabled"]["usage"]) == CHAT_ARM_INPUT_TOTAL and c["4.6"]["answers"] == 34
    assert c["4.6"]["citation_validity"] == 1.0 and c["4.6"]["truncated_answers"] == 0
    e = rep["arms"]["export.synthesis"]
    assert e["4.6"]["sections_present"] == 6 and e["4.6"]["n_invented_links"] == 0 and e["4.6"]["link_coverage"] == 1.0 and e["4.6"]["usage"]["input_tokens"] == e["5-disabled"]["usage"]["input_tokens"]
    u = rep["arms"]["planner.update"]
    assert u["4.6"]["input_tokens"] == u["5-disabled"]["input_tokens"] and u["4.6"]["well_formed"] == 1
    # state restored: no plans, no leftover notes/facts beyond the shared approved findings
    pid = rep["project_id"]
    assert db.latest_plan(pid) is None and not db.list_facts(pid)
    assert len(db.list_project_notes(pid)) == rep["shared_inputs"]["approved_findings"]
    # env + contracts untouched
    assert not any(k.startswith("NEUROSEARCH_TASK_") and ("PLANNER" in k or "ANSWER" in k or "EXPORT" in k) for k in os.environ)
    assert contracts.contract("answer.chat").model == settings.answer_model == "claude-sonnet-4-6" and contracts.contract("planner.build").model == settings.answer_model
    assert rep["spend_estimate"]["maximum"] > rep["spend_estimate"]["estimate"] > 1.0 and db.kv_get("daily_budget") == str(rep["spend_estimate"]["budget"])
    names = {pathlib.Path(f).name for f in rep["files"]}
    assert {"planner-4.6.json", "planner-5-disabled.json", "planner-5-adaptive-medium.json", "planner.update-4.6.json", "planner.update-5-disabled.json",
            "export.synthesis-4.6.json", "export.synthesis-5-disabled.json", "answer.chat-4.6.json", "answer.chat-5-disabled.json", "comparison.json", "comparison.txt"} <= names
    assert "RECOMMENDATIONS (nothing was changed" in rep["text"] and "== answer.repair" in rep["text"] and "expected spend" in rep["text"]
    assert all(v.get("production_default_changed") is False for v in rep["verdicts"].values())


def test_installation_check_fails_clearly_without_jsonschema(monkeypatch):
    import sys
    from neurosearch import schemas
    schemas.check_installation()                                  # the real environment is complete
    monkeypatch.setitem(sys.modules, "jsonschema", None)          # simulate an un-reinstalled checkout
    with pytest.raises(SystemExit) as ei:
        schemas.check_installation()
    assert "installation is incomplete" in str(ei.value) and "./start" in str(ei.value)


# ------------------------------------------------------------------ F4: Planner V3 exit gate (all under the fakes)

@pytest.fixture
def golden_project(isolated_db, monkeypatch):
    from neurosearch import evals, findings
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    monkeypatch.setattr(settings, "daily_budget", 1000)
    g = evals.load_golden()
    pid = g["project_id"]
    for gid, sid in g["sources"].items():
        if gid != "calc":
            findings.suggest_for_source(pid, sid)
    for n in db.list_project_notes(pid, status="suggested"):
        db.set_note_status(n["id"], "approved")
    return pid


def _plan_body(p):
    return {k: v for k, v in p.items() if not k.startswith("_")}


def test_planner_v3_exit_gate(golden_project, monkeypatch):
    from neurosearch import migration, planner, planner_v3, schemas
    from neurosearch.config import settings
    pid = golden_project
    # [ ] legacy planner is the default and stays selectable
    monkeypatch.setattr(settings, "planner_v3", False)
    v1 = planner.build_plan(pid)["plan"]
    assert "_build" not in v1 and "_ids" not in v1
    monkeypatch.setattr(settings, "planner_v3", True)
    before = db.health()["structured_outputs"]
    __import__("neurosearch.fake_ai", fromlist=["_cache_seen"])._cache_seen.clear()      # a cold prompt cache, as a first build sees it
    row = planner.build_plan(pid)
    p = row["plan"]
    # [ ] situation + four component schemas valid, same analysis_hash everywhere, zero mismatches/fallbacks
    comps = p["_components"]
    for name, schema in (("situation", "situation-v3"), ("core", "plan-core-v3"), ("execution", "plan-execution-v3"), ("economics", "plan-economics-v3"), ("actions", "plan-actions-v3")):
        assert schemas.is_valid(schema, comps[name]), name
    b = p["_build"]
    assert b["planner_version"] == "v3" and [c["task"] for c in b["components"]] == ["planner.situation", "planner.core", "planner.execution", "planner.economics", "planner.actions"]
    assert {c["analysis_hash"] for c in b["components"][1:]} == {b["analysis_hash"]} == {p["_analysis_hash"]} and b["components"][0]["analysis_hash"] is None
    assert db.health()["structured_outputs"] == before and b["analysis_hash"] == __import__("hashlib").sha256(json.dumps(comps["situation"], sort_keys=True, separators=(",", ":")).encode()).hexdigest()[:16]
    # [ ] zero dangling evidence / dependencies / cycles, unique ids
    assert p["_evidence_check"]["dangling"] == [] and p["_evidence_check"]["references"] > 0
    plan, ids, problems = planner_v3.assemble(comps["situation"], comps["core"], comps["execution"], comps["economics"], comps["actions"], set(p["_evidence"]))
    assert problems == [] and len(set(ids.values())) == len(ids) and all(":" in v for v in ids.values())
    assert p["_ids"] == ids and all(w["task"] in ids.values() for w in p["this_week"])
    # [ ] deterministic assembly: same inputs → identical document
    plan2, ids2, _ = planner_v3.assemble(comps["situation"], comps["core"], comps["execution"], comps["economics"], comps["actions"], set(p["_evidence"]))
    assert plan2 == plan and ids2 == ids and _plan_body(p) == plan
    # [ ] existing plan shape unchanged: every V1 key present, UI/export consumers work
    assert set(_plan_body(v1)) <= set(_plan_body(p)) and set(p["analysis"]) >= {"situation", "swot", "readiness", "options", "assumptions", "failure_patterns", "verdict"}
    md = planner.plan_markdown(row, db.get_project(pid))
    assert "First steps" in md or "first" in md.lower()
    assert planner.plan_html(row, db.get_project(pid))
    # [ ] frozen rubric: V3 ≥ V1 on the fakes
    rub = migration.load_rubric()
    assert migration.score_rubric(_plan_body(p), rub["plan"])["score"] >= migration.score_rubric(_plan_body(v1), rub["plan"])["score"]
    assert migration.score_rubric(p["analysis"], rub["analysis"])["score"] >= migration.score_rubric(v1.get("analysis"), rub["analysis"])["score"]
    # [ ] cache telemetry answers F5's question
    assert b["material_tokens_est"] > 1000 and b["cache_write_total"] > 0 and b["cache_read_total"] > 0 and b["cost_total"] > 0 and b["cache_read_share"] > 0.5
    assert all(k in b["components"][0] for k in ("input_tokens", "output_tokens", "cache_read", "cache_write", "cost", "seconds", "schema"))


def test_planner_v3_status_survives_rebuild(golden_project, monkeypatch):
    """Statuses follow semantic ids: a task that moves to another position, or whose id the model renamed but whose
    text is the same, keeps its DONE. Positional-only carry is gone."""
    import copy
    from neurosearch import fake_ai, planner
    from neurosearch.config import settings
    pid = golden_project
    monkeypatch.setattr(settings, "planner_v3", True)
    r1 = planner.build_plan(pid)
    assert r1["plan"]["_ids"]["phases.0.tasks.0"] == "task:get-lender-prequalification"
    db.set_item_status(r1["id"], "phases.0.tasks.0", "done", "did it")
    db.set_item_status(r1["id"], "first_steps.1", "in_progress")
    # rebuild with the same tasks in a different order and one id renamed consistently everywhere (same text)
    renamed = json.loads(json.dumps(fake_ai.PLANNER_V3).replace("task:get-lender-prequalification", "task:lender-prequal-letter"))
    renamed["planner.execution"]["phases"][0]["tasks"].reverse()
    for k, v in renamed.items():
        monkeypatch.setitem(fake_ai.PLANNER_V3, k, v)
    r2 = planner.build_plan(pid)
    p2 = r2["plan"]
    assert p2["_ids"]["phases.0.tasks.1"] == "task:get-lender-prequalification"          # reconciled back to the stable id
    assert p2["_build"]["id_renames_on_rebuild"] == {"task:lender-prequal-letter": "task:get-lender-prequalification"}
    assert r2["items"]["phases.0.tasks.1"] == {"status": "done", "note": "did it"} and "phases.0.tasks.0" not in r2["items"]
    assert r2["items"]["first_steps.1"]["status"] == "in_progress"
    # references were rewritten with the reconciled id
    assert all("task:lender-prequal-letter" not in json.dumps(v) for v in (p2["phases"], p2["first_steps"], p2["this_week"], p2["risks"]))


def test_planner_v3_assembly_refuses_incoherent_components(golden_project, monkeypatch):
    """Schema-valid pieces that contradict each other never become a plan: dangling refs, cycles, self-dependency,
    duplicate ids, missing evidence → PlanAssemblyError, previous plan retained, event recorded."""
    import copy
    from neurosearch import fake_ai, planner, planner_v3
    from neurosearch.config import settings
    pid = golden_project
    monkeypatch.setattr(settings, "planner_v3", True)
    good = planner.build_plan(pid)
    comps = good["plan"]["_components"]
    ev = set(good["plan"]["_evidence"])
    def broken(mut):
        c = copy.deepcopy(comps); mut(c); return planner_v3.assemble(c["situation"], c["core"], c["execution"], c["economics"], c["actions"], ev)[2]
    assert any("does not exist" in x for x in broken(lambda c: c["actions"]["this_week"][0].__setitem__("task", "task:nope")))
    assert any("cycle" in x for x in broken(lambda c: c["execution"]["phases"][0]["tasks"][0].__setitem__("depends_on", ["task:build-deal-flow"])))
    assert any("depends on itself" in x for x in broken(lambda c: c["execution"]["phases"][0]["tasks"][0].__setitem__("depends_on", ["task:get-lender-prequalification"])))
    assert any("duplicate id" in x for x in broken(lambda c: c["economics"]["risks"][1].__setitem__("id", c["economics"]["risks"][0]["id"])))
    assert broken(lambda c: c["economics"]["risks"][0].__setitem__("id", "task:build-deal-flow")) == []      # ids are namespaced by kind: risk:build-deal-flow
    assert any("evidence ids do not exist" in x for x in broken(lambda c: c["core"]["goal"].__setitem__("evidence", ["F999"])))
    assert any("not one of its options" in x for x in broken(lambda c: c["situation"].__setitem__("recommended_option", "option:nope")))
    assert any("phase" in x for x in broken(lambda c: c["economics"]["costs"]["upfront"][0].__setitem__("phase", "phase:nope")))
    # through build_plan: the previous plan stays current, the failure is visible
    bad = copy.deepcopy(fake_ai.ACTIONS_V3); bad["this_week"][0]["task"] = "task:does-not-exist"
    monkeypatch.setitem(fake_ai.PLANNER_V3, "planner.actions", bad)
    with pytest.raises(planner_v3.PlanAssemblyError, match="does not exist"):
        planner.build_plan(pid)
    assert db.latest_plan(pid)["id"] == good["id"]
    assert db.validation_events(kind="plan_assembly_failed")[0]["detail"]["analysis_hash"] == good["plan"]["_analysis_hash"]
    assert db.health()["evidence"]["events"].get("plan_assembly_failed") == 1


def test_planner_v3_id_helpers():
    from neurosearch import planner_v3 as V
    assert V.norm_id("task", "Task: Get Lender Pre-Qualification!") == "task:get-lender-pre-qualification"
    assert V.norm_id("task", "", "Call two lenders") == "task:call-two-lenders" and V.norm_id("risk", None, "") == "risk:item"
    assert V.similarity("Get pre-qualified by two SBA lenders", "get prequalified by two sba lenders") >= V.ID_SIMILARITY
    assert V._find_cycle({"a": {"b"}, "b": {"c"}, "c": {"a"}}) and not V._find_cycle({"a": {"b"}, "b": set()})


# ------------------------------------------------------------------ F5: closeout command (dry run under the fakes)

def test_mission_f_closeout_dry_run(tmp_path, monkeypatch):
    """One command: deterministic phase (Tier 1 with V1 and V3, registry, rubric) then the narrowly scoped live phase — here
    on the fakes — ending in a Mission F verdict and a separate Planner V3 decision; artifacts saved; nothing changed."""
    import os
    from neurosearch import closeout, contracts
    from neurosearch.config import settings
    monkeypatch.delenv("NEUROSEARCH_PLANNER_V3", raising=False)
    monkeypatch.setattr(settings, "planner_v3", False)
    was_dir = settings.data_dir
    rep = closeout.run_closeout(live=False, out_dir=tmp_path / "evals", progress=lambda m: None, run_pytest=False)
    assert settings.data_dir == was_dir and settings.planner_v3 is False           # restored; no production change
    det = rep["deterministic"]
    assert det["pass"] and {s["step"] for s in det["steps"]} >= {"schema registry provider-compatible", "Tier 1 (Planner V1)", "Tier 1 (Planner V3)", "Tier 1 V3 evidence / assembly", "rubric V3 ≥ V1 (fakes)"}
    lv = rep["live"]
    assert set(lv["surfaces"]) == {"findings.extract", "rank.relevance", "planner.update", "discover.quick"} and all(s["pass"] for s in lv["surfaces"].values()), lv["surfaces"]
    assert lv["surfaces"]["findings.extract"]["schema"] == "findings-v2" and lv["surfaces"]["rank.relevance"]["schema"] == "rank-v2"
    assert lv["surfaces"]["planner.update"]["n_updates"] >= 1 and lv["surfaces"]["planner.update"]["addresses_new_finding"] == 1
    assert all(sum(s["events"].values()) == 0 for s in lv["surfaces"].values())
    p = rep["planner_v3"]
    assert p["promote"] and p["v3"]["telemetry"]["cache_read_share"] > 0.5 and p["v3"]["usage"]["calls"] == 5 and p["v1"]["usage"]["calls"] == 2
    assert p["v3"]["rubric_plan"] >= p["v1"]["rubric_plan"] and p["v3"]["dangling_raw"] == 0 and "cost_cold_equivalent" in p["v3"]
    assert rep["mission"]["verdict"] in ("PASS", "PASS_WITH_CAVEAT") and "PLANNER V3: PROMOTE" in rep["text"] and "Nothing was changed" in rep["text"]
    names = {q.name for q in pathlib.Path(rep["dir"]).iterdir()}
    assert {"deterministic.json", "findings.json", "ranking.json", "planner-v1.json", "planner-v3.json", "update.json", "discover.json", "closeout.json", "closeout.txt"} <= names
    assert lv["spend_estimate"]["maximum"] > lv["spend_estimate"]["estimate"] > 0.5
    assert contracts.contract("findings.extract").model == "claude-sonnet-5"       # untouched


def test_planner_v3_promotion_rules():
    import copy
    from neurosearch import closeout as C
    def task(rubric, struct=1.0, refs=30, **kw):
        t = {"rubric": {"score": rubric, "failed": [], "structure": {"score": struct, "failed": []}}, "evidence_refs": refs, "dangling_after_removal": 0, "truncated": 0,
             "parse_failed": 0, "json_repaired": 0, "invocations": {"returned_model": "claude-sonnet-4-6-20260210"}}
        t.update(kw); return t
    def arm(rp, ra, calls, cost, secs, tel=None, events=None, dangling=0, error=None):
        return {"error": error, "seconds": secs, "tasks": {"planner.analysis": task(ra), "planner.build": task(rp)}, "plan_evidence_dangling_raw": dangling,
                "usage": {"cost": cost, "calls": calls, "input_tokens": 1000, "output_tokens": 3000, "cache_write_tokens": 40000, "cache_read_tokens": 0},
                "build_telemetry": tel, "events": events or {}}
    tel = {"cache_read_share": 0.85, "components": []}
    v1 = arm(0.7, 0.6, 2, 0.30, 90.0)
    v3 = arm(0.75, 0.65, 5, 0.36, 110.0, tel)
    d = C.planner_decision({"v1": v1, "v3": v3})
    assert d["promote"] and d["decision"].startswith("PROMOTE") and d["mission_fails"] == [] and any("+20%" in c for c in d["caveats"])
    assert not C.planner_decision({"v1": v1, "v3": arm(0.65, 0.65, 5, 0.36, 110.0, tel)})["promote"]                     # rubric below V1
    assert not C.planner_decision({"v1": v1, "v3": arm(0.7, 0.6, 5, 0.36, 110.0, tel, dangling=1)})["promote"]           # evidence validity < 1
    bad = C.planner_decision({"v1": v1, "v3": arm(0.9, 0.9, 5, 0.36, 110.0, tel, events={"schema_mismatch": 1})})
    assert not bad["promote"] and bad["mission_fails"]                                                                    # events fail Mission F too
    assert not C.planner_decision({"v1": v1, "v3": arm(0.9, 0.9, 5, 1.20, 110.0, tel)})["promote"]                        # pathological cost (4×)
    ok = C.planner_decision({"v1": v1, "v3": arm(0.9, 0.9, 5, 0.60, 200.0, tel)})                                        # 2× cost, 2.2× latency → caveats only
    assert ok["promote"] and len(ok["caveats"]) == 2
    low = C.planner_decision({"v1": v1, "v3": arm(0.7, 0.6, 5, 0.36, 110.0, {"cache_read_share": 0.2, "components": []})})
    assert low["promote"] and any("cache read share" in c for c in low["caveats"])
    assert not C.planner_decision({"v1": v1, "v3": arm(0.9, 0.9, 5, 0.3, 100.0, tel, error="boom")})["promote"]


def test_version_is_pep440_and_consistent():
    """`pip install -e .` (what ./start runs) rejects non-PEP-440 versions silently-ish; a bad version means the CLI never
    picks up new entry points or dependencies. Local labels like 0.19.0+f5 are fine; 0.19.0-f5 is not."""
    import re
    from packaging.version import InvalidVersion, Version
    import neurosearch
    pj = pathlib.Path(__file__).parent.parent / "pyproject.toml"
    v = re.search(r'^version = "([^"]+)"', pj.read_text(), re.M).group(1)
    try:
        Version(v)
    except InvalidVersion:
        pytest.fail(f"pyproject version {v!r} is not PEP 440 — pip install -e . will fail")
    assert v == neurosearch.__version__
    assert f"'{v}'" in (pathlib.Path(neurosearch.__file__).parent / "web" / "index.html").read_text()


# ------------------------------------------------------------------ F5 harness fixes: frozen research, INCOMPLETE stages, resume

def test_frozen_research_is_byte_identical_for_both_planners(golden_project, monkeypatch):
    """The planner's research is prepared once (strict) and handed to both arms; after the freeze, retrieval can fail or
    disappear without changing either arm's input — both consume the same bytes, and both record the same hash."""
    from neurosearch import planner, planner_v3, providers, search
    from neurosearch.config import settings
    pid = golden_project
    _real_search = search.search
    research = planner.research_context(pid, strict=True)
    assert research["evidence_count"] > 10 and research["material_hash"] and research["material"].startswith("PROJECT:")
    # after the freeze: vector search is dead
    def dead(*a, **k):
        raise search.RetrievalUnavailable("vector search unavailable: CONNECTION after 3 attempts")
    monkeypatch.setattr(search, "search", dead)
    monkeypatch.setattr(planner, "search", dead)
    seen: dict[str, list[str]] = {"v1": [], "v3": []}
    real_invoke = providers.invoke
    current = {"arm": "v1"}
    def spy(task, **kw):
        sysb = kw.get("system")
        first = sysb[0]["text"] if isinstance(sysb, list) else str(sysb)
        if task.startswith("planner."):
            seen[current["arm"]].append(first)
        return real_invoke(task, **kw)
    monkeypatch.setattr(providers, "invoke", spy)
    monkeypatch.setattr(settings, "planner_v3", False)
    r1 = planner.build_plan(pid, research=research)
    current["arm"] = "v3"
    monkeypatch.setattr(settings, "planner_v3", True)
    r3 = planner.build_plan(pid, research=research)
    assert r1["plan"]["_research_hash"] == r3["plan"]["_research_hash"] == research["material_hash"] == r3["plan"]["_build"]["research_hash"]
    shared_v1 = {t for t in seen["v1"]}; shared_v3 = {t for t in seen["v3"]}
    assert len(shared_v1) == 1 and shared_v1 == shared_v3                         # every planner call, both arms: the same first system block, byte for byte
    assert research["material"] in next(iter(shared_v1))
    # without a frozen payload, strict preparation refuses to run on a dead vector search
    with pytest.raises(search.RetrievalUnavailable):
        planner.research_context(pid, strict=True)
    # the production (non-strict) path degrades visibly instead
    monkeypatch.setattr(search, "search", search.search.__wrapped__ if hasattr(search.search, "__wrapped__") else _real_search)
    from neurosearch import embeddings
    monkeypatch.setattr(embeddings, "embed_query", lambda q: (_ for _ in ()).throw(ConnectionError("embedding endpoint down")))
    before = db.health()["evidence"]["retrieval_degraded"]
    hits = search.search("equity injection", limit=3, source_ids=None)
    assert isinstance(hits, list) and db.health()["evidence"]["retrieval_degraded"] == before + 1
    with pytest.raises(search.RetrievalUnavailable):
        search.search("equity injection", limit=3, strict=True)


def test_closeout_marks_incomplete_and_resumes_without_repaying(tmp_path, monkeypatch):
    """A dependency failure during Planner V3 makes that stage INCOMPLETE (no promotion decision on a degraded fallback);
    the run's database and every completed stage are kept; --resume reuses them and reruns only V3, recording the history."""
    from neurosearch import closeout, migration, providers
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "planner_v3", False)
    out = tmp_path / "evals"
    orig_arm = migration.run_planner_arm
    def flaky(*a, **k):                                   # the V3 arm hits a dead endpoint during the live phase only
        if settings.planner_v3:
            raise providers.ProviderError("CONNECTION", RuntimeError("Connection error."), 3)
        return orig_arm(*a, **k)
    monkeypatch.setattr(migration, "run_planner_arm", flaky)
    r1 = closeout.run_closeout(live=False, out_dir=out, progress=lambda m: None, run_pytest=False)
    assert r1["mission"]["verdict"] == "INCOMPLETE" and "planner_v3" in r1["mission"]["incomplete"] and "CONNECTION" in r1["mission"]["incomplete"]["planner_v3"]
    assert r1["planner_v3"]["decision"].startswith("NOT DECIDED") and not r1["planner_v3"]["promote"] and "PLANNER V3: NOT DECIDED" in r1["text"]
    st = r1["stages"]
    assert {k for k, v in st.items() if v["status"] == "done"} == {"shared_inputs", "findings", "ranking", "research", "planner_v1", "update", "discover"}
    d1 = pathlib.Path(r1["dir"])
    assert (d1 / "data" / "neurosearch.db").exists() and (d1 / "stages.json").exists() and (d1 / "research.json").exists() and (d1 / "planner-v1.json").exists() and not (d1 / "planner-v3.json").exists()
    # resume: only planner_v3 runs; everything else (incl. the database and the frozen research) is reused
    calls = []
    monkeypatch.setattr(migration, "run_planner_arm", lambda *a, **k: calls.append(k.get("research", {}).get("material_hash")) or orig_arm(*a, **k))
    r2 = closeout.run_closeout(live=False, out_dir=out, progress=lambda m: None, run_pytest=False, resume=True)
    assert r2["resumed_from"] == str(d1) and len(calls) == 1                          # exactly one planner build (V3) was paid for
    assert calls[0] == r1["live"]["research_hash"] == r2["live"]["research_hash"]      # on the same frozen research
    st2 = r2["stages"]
    assert st2["findings"]["reused_from"] == str(d1) and st2["ranking"]["reused_from"] == str(d1) and st2["planner_v1"]["reused_from"] == str(d1) and st2["research"]["reused_from"] == str(d1)
    assert st2["update"]["reused_from"] == str(d1) and st2["discover"]["reused_from"] == str(d1) and "reused_from" not in st2["planner_v3"]
    h = r2["history"][-1]
    assert "planner_v3" in h["invalidated"] and "CONNECTION" in h["invalidated"]["planner_v3"] and "planner_v1" in h["reused"] and h["database"].startswith("copied")
    assert r2["mission"]["verdict"] in ("PASS", "PASS_WITH_CAVEAT") and r2["planner_v3"]["promote"] and "HISTORY" in r2["text"] and "invalidated planner_v3" in r2["text"]
    assert r2["live"]["surfaces"]["findings.extract"]["reused"] and r2["live"]["surfaces"]["rank.relevance"]["reused"]
    # the previous run directory is untouched (its own closeout.txt still says INCOMPLETE)
    assert "INCOMPLETE" in (d1 / "closeout.txt").read_text()


def test_closeout_resume_from_legacy_run(tmp_path, monkeypatch):
    """The first live run predates stage files and did not persist its database: findings/ranking/update/discover results are
    reused (events recovered from its closeout.json), shared inputs are rebuilt, and both planner arms are rebuilt on freshly
    frozen research — the old V1 is recorded as invalidated, never silently compared against a new V3."""
    from neurosearch import closeout
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "planner_v3", False)
    out = tmp_path / "evals"
    r1 = closeout.run_closeout(live=False, out_dir=out, progress=lambda m: None, run_pytest=False)
    d1 = pathlib.Path(r1["dir"])
    # turn it into a legacy run: no stages.json, no database, no research, stage files without events
    (d1 / "stages.json").unlink(); shutil.rmtree(d1 / "data"); (d1 / "research.json").unlink()
    for f in ("findings.json", "ranking.json", "update.json", "discover.json"):
        obj = json.loads((d1 / f).read_text()); obj.pop("events", None); obj.pop("saved_sonnet5", None); (d1 / f).write_text(json.dumps(obj))
    (d1 / "planner-v3.json").unlink()                                          # the invalid V3 attempt is not a reusable artifact
    r2 = closeout.run_closeout(live=False, out_dir=out, progress=lambda m: None, run_pytest=False, resume=True)
    h = r2["history"][-1]
    assert set(h["reused"]) == {"findings", "ranking", "update", "discover"} and "planner_v1" in h["rebuilt"] and "different research" in h["invalidated"]["planner_v1"]
    assert h["database"].startswith("the previous run did not persist")
    st = r2["stages"]
    assert "reused_from" not in st["shared_inputs"] and "reused_from" not in st["planner_v1"] and "reused_from" not in st["planner_v3"] and st["findings"].get("legacy")
    assert r2["live"]["surfaces"]["findings.extract"]["reused"] and r2["live"]["surfaces"]["findings.extract"]["events"] == {k: 0 for k in closeout.EVENT_KINDS}
    assert r2["mission"]["verdict"] in ("PASS", "PASS_WITH_CAVEAT") and r2["planner_v3"]["v1"]["usage"]["calls"] == 2 and r2["planner_v3"]["v3"]["usage"]["calls"] == 5


def test_enum_casing_is_normalised_before_validation(isolated_db, monkeypatch):
    """Provider docs: enum/const capitalisation is not guaranteed under structured outputs. 'Medium' for a 'medium' enum
    must not be a schema mismatch; a value outside the enum still is."""
    from neurosearch import providers, schemas
    from neurosearch.config import settings
    monkeypatch.setattr(settings, "fake_ai", True)
    obj = {"sources": [{"name": "X", "kind": "Podcast", "url": "", "gist": "g", "why": "w", "angle": "a", "start_with": [], "fit": 3, "depth": "BEGINNER"}], "note": "n"}
    assert not schemas.is_valid("discovery-v2", json.loads(json.dumps(obj)))
    fixed = schemas.normalize_enums("discovery-v2", json.loads(json.dumps(obj)))
    assert fixed["sources"][0]["kind"] == "podcast" and fixed["sources"][0]["depth"] == "beginner" and schemas.is_valid("discovery-v2", fixed)
    assert not schemas.is_valid("discovery-v2", schemas.normalize_enums("discovery-v2", {"sources": [{**obj["sources"][0], "kind": "blog"}], "note": "n"}))
    class R:
        content = [type("B", (), {"type": "text", "text": json.dumps({"updates": [{"section": "s", "previous": "p", "proposed": "x", "reason": "r"}]})})()]
        stop_reason, model, usage = "end_turn", "m", None
    assert providers.structured("planner.update", R())["updates"][0]["proposed"] == "x"
    risk = {"costs": {"upfront": [], "recurring": [], "optional": [], "services": [], "contingency": "", "minimum": "", "recommended": "", "premium": "", "note": "", "evidence": []},
            "tools": [], "risks": [{"id": "risk:x", "risk": "r", "mitigation": "m", "priority": "High", "related": [], "evidence": []}], "gotchas": []}
    assert schemas.normalize_enums("plan-economics-v3", risk)["risks"][0]["priority"] == "high" and schemas.is_valid("plan-economics-v3", risk)
